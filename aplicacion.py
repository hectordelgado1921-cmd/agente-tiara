import json
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import streamlit as st

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
DATA.mkdir(exist_ok=True)
KB_FILE = DATA / "base_conocimiento_tiara.json"
DEFAULT_EXCEL = next(DATA.glob("*.xlsx"), None)
DEFAULT_JSON = DATA / "base_conocimiento_tiara_septiembre_2026.json"


def clean(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return re.sub(r"\s+", " ", v.strip()) if isinstance(v, str) else v


def norm(v):
    s = "" if v is None else str(v).lower()
    return s.translate(str.maketrans("áéíóúüñ", "aeiouun"))


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return None
    m = re.search(r"-?\d+(?:[.,]\d+)?", str(v).replace(",", ""))
    return float(m.group()) if m else None


def header_row(ws):
    best = (1, -1)
    for r in range(1, min(20, ws.max_row) + 1):
        score = sum(
            clean(ws.cell(r, c).value) not in (None, "")
            for c in range(1, ws.max_column + 1)
        )
        if score > best[1]:
            best = (r, score)
    return best[0]


def read_sheet(ws):
    hr = header_row(ws)
    headers = {}
    for c in range(1, ws.max_column + 1):
        v = clean(ws.cell(hr, c).value)
        if v not in (None, ""):
            headers[c] = str(v)

    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        d = {}
        for c, h in headers.items():
            v = clean(ws.cell(r, c).value)
            if v not in (None, ""):
                d[h] = v
        if d:
            rows.append({"source_row": r, "data": d})
    return {"header_row": hr, "records": rows}


def classify(text):
    t = norm(text)
    if "garantia" in t:
        return "GARANTIA"
    if "instal" in t:
        return "INSTALACION"
    if "mantenimiento" in t or "aceite" in t or "filtro" in t:
        return "MANTENIMIENTO"
    if "repar" in t or "arreglo" in t:
        return "REPARACION"
    if any(x in t for x in ("falla", "calent", "problema", "no funciona", "no trabaja")):
        return "FALLA"
    return "OBSERVACION"


def import_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    kb = {
        "metadata": {
            "agent": "Agente Tiara",
            "version": "1.0",
            "source_of_truth": "Excel",
            "source_file": Path(path).name,
            "updated_at": datetime.now().isoformat(timespec="seconds"),
        },
        "sheets": {},
        "maintenance_records": [],
        "log_records": [],
        "inventory_items": [],
        "fuel_records": [],
        "budget_records": [],
        "invoice_records": [],
        "cleaning_records": [],
        "documents": [],
    }

    for name in wb.sheetnames:
        block = read_sheet(wb[name])
        kb["sheets"][name] = block
        n = norm(name)
        group = None

        if "checklist" in n:
            group = "maintenance_records"
        elif "vitacora" in n or "bitacora" in n:
            group = "log_records"
            for rec in block["records"]:
                rec["event_type"] = classify(
                    " | ".join(map(str, rec["data"].values()))
                )
        elif n.strip() == "inventario":
            group = "inventory_items"
        elif "combustible" in n:
            group = "fuel_records"
        elif "presupuesto" in n:
            group = "budget_records"
        elif "facturas" in n:
            group = "invoice_records"
        elif "limpieza" in n:
            group = "cleaning_records"
        elif "permisos" in n or "seguros" in n:
            group = "documents"

        if group:
            kb[group].extend([{"source_sheet": name, **r} for r in block["records"]])

    save(kb)
    return kb


def save(kb):
    KB_FILE.write_text(
        json.dumps(kb, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )


def load():
    return json.loads(KB_FILE.read_text(encoding="utf-8")) if KB_FILE.exists() else None


def all_records(kb):
    for g in (
        "maintenance_records",
        "log_records",
        "inventory_items",
        "fuel_records",
        "budget_records",
        "invoice_records",
        "cleaning_records",
        "documents",
    ):
        for r in kb[g]:
            yield g, r


def search(kb, q, groups=None, limit=20):
    terms = [x for x in re.findall(r"[a-z0-9]+", norm(q)) if len(x) > 2]
    out = []
    for g, r in all_records(kb):
        if groups and g not in groups:
            continue
        blob = norm(json.dumps(r, ensure_ascii=False))
        score = sum(blob.count(t) for t in terms)
        if score:
            out.append((score, g, r))
    out.sort(key=lambda x: x[0], reverse=True)
    return out[:limit]


def costs(kb, q):
    terms = [x for x in re.findall(r"[a-z0-9]+", norm(q)) if len(x) > 2]
    total = 0
    rows = []
    for g, r in all_records(kb):
        d = r["data"]
        blob = norm(json.dumps(d, ensure_ascii=False))
        if terms and not any(t in blob for t in terms):
            continue
        for k, v in d.items():
            if "total" in norm(k):
                n = num(v)
                if n is not None:
                    total += n
                    rows.append((g, r["source_row"], k, n))
    return total, rows


def choose(q):
    t = norm(q)
    if any(x in t for x in (
        "mantenimiento", "aceite", "filtro", "anodo", "bujia",
        "impeler", "impeller", "servicio"
    )):
        return "mantenimiento"
    if any(x in t for x in (
        "falla", "problema", "calent", "temperatura", "bitacora", "repar"
    )):
        return "eventos"
    if any(x in t for x in ("inventario", "stock", "existencia", "repuesto")):
        return "inventario"
    if any(x in t for x in ("factura", "proveedor")):
        return "facturas"
    if any(x in t for x in ("combustible", "gasolina", "diesel", "litros")):
        return "combustible"
    if any(x in t for x in (
        "permiso", "seguro", "vence", "vencimiento", "documento"
    )):
        return "documentos"
    if any(x in t for x in ("cuanto gaste", "cuánto gasté", "costo", "costos", "presupuesto")):
        return "costos"
    return "historial"


def parse_maintenance_date(value):
    if value is None:
        return []
    text = str(value)
    found = []
    for m in re.finditer(r"(?<!\d)(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})(?!\d)", text):
        day, month, year = map(int, m.groups())
        if year < 100:
            year += 2000
        try:
            found.append((date(year, month, day), m.group(0)))
        except ValueError:
            pass
    for m in re.finditer(r"(?<!\d)(\d{1,2})-(\d{2})(\d{2})(?!\d)", text):
        day, month, year = map(int, m.groups())
        year += 2000
        try:
            found.append((date(year, month, day), m.group(0)))
        except ValueError:
            pass
    return found


def latest_maintenance(kb, q):
    t = norm(q)
    if not any(x in t for x in ("ultimo", "ultima", "ultimos", "ultimas", "reciente", "recientes")):
        return None

    records = kb.get("maintenance_records", [])
    target_motors = any(x in t for x in ("motor", "motores")) and "generador" not in t
    target_kit = any(x in t for x in ("kit 300", "300 horas", "300h"))
    target_patas = any(x in t for x in ("pata", "patas"))
    target_generator = "generador" in t

    if target_generator:
        selected = []
        in_generator = False
        for r in records:
            name = norm(r.get("data", {}).get("mantenimiento", "")).strip()
            if name in ("generador", "generador electrico"):
                in_generator = True
                continue
            if in_generator:
                selected.append(r)
        if not selected:
            selected = [r for r in records if "generador" in norm(json.dumps(r.get("data", {}), ensure_ascii=False))]
    elif target_kit:
        selected = [
            r for r in records
            if "kit 300 horas" in norm(r.get("data", {}).get("mantenimiento", ""))
        ]
    elif target_patas:
        selected = [
            r for r in records
            if "pata" in norm(r.get("data", {}).get("mantenimiento", ""))
        ]
    elif target_motors:
        selected = []
        in_motors = False
        for r in records:
            name = norm(r.get("data", {}).get("mantenimiento", "")).strip()
            if name in ("motores", "motor"):
                in_motors = True
                continue
            if in_motors and name in ("generador", "generador electrico"):
                break
            if in_motors:
                selected.append(r)
        if not selected:
            selected = [r for r in records if "motor" in norm(json.dumps(r.get("data", {}), ensure_ascii=False))]
    else:
        selected = [r for r in records if norm(r.get("data", {}).get("mantenimiento", "")).strip() not in ("motores", "motor", "generador", "generador electrico")]

    candidates = []
    for r in selected:
        for key, value in r.get("data", {}).items():
            if "fecha" not in norm(key):
                continue
            for dt, raw in parse_maintenance_date(value):
                candidates.append((dt, r, raw))

    if not candidates:
        return None

    latest_date = max(x[0] for x in candidates)
    latest_rows = []
    seen = set()
    for dt, r, raw in candidates:
        if dt == latest_date and r["source_row"] not in seen:
            latest_rows.append(r)
            seen.add(r["source_row"])
    return latest_date, latest_rows


def render(results):
    if not results:
        return "No encontré registros relacionados en la base."
    lines = []
    for score, g, r in results[:10]:
        text = " | ".join(
            f"{k}: {v}" for k, v in r["data"].items() if v not in (None, "")
        )
        lines.append(f"**[{g}]** · fila {r['source_row']}\n\n{text[:1000]}")
    return "\n\n---\n\n".join(lines)


def answer(kb, q):
    t = norm(q)
    latest = latest_maintenance(kb, q)
    if latest is not None:
        latest_date, rows = latest
        if any(x in t for x in ("motor", "motores")) and "generador" not in t:
            label = "motores"
        elif any(x in t for x in ("kit 300", "300 horas", "300h")):
            label = "kit de 300 horas"
        elif any(x in t for x in ("pata", "patas")):
            label = "patas de motores"
        elif "generador" in t:
            label = "generador"
        else:
            label = "mantenimiento"
        fecha = f"{latest_date.day}-{latest_date.month}-{str(latest_date.year)[-2:]}"
        return f"### Último mantenimiento de {label}\n\n**{fecha}**"

    c = choose(q)
    groups = {
        "mantenimiento": ["maintenance_records"],
        "eventos": ["log_records"],
        "inventario": ["inventory_items"],
        "facturas": ["invoice_records"],
        "combustible": ["fuel_records"],
        "documentos": ["documents"],
        "historial": None,
    }.get(c)

    if c == "costos":
        total, rows = costs(kb, q)
        return f"### Resultado de costos\n**Total: ${total:,.2f} USD**\n\nRegistros encontrados: **{len(rows)}**"

    return f"### Resultados: {c.title()}\n\n" + render(search(kb, q, groups))

def load_default():
    if KB_FILE.exists():
        return load()
    if DEFAULT_JSON.exists():
        return json.loads(DEFAULT_JSON.read_text(encoding="utf-8"))
    if DEFAULT_EXCEL and DEFAULT_EXCEL.exists():
        return import_excel(DEFAULT_EXCEL)
    return None


st.set_page_config(
    page_title="Agente Tiara",
    page_icon="⚓",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
    .stApp {
        background: #0e1117 !important;
        color: #f1f5f9 !important;
    }
    [data-testid="stHeader"] {
        background: rgba(14,17,23,0);
    }
    [data-testid="stSidebar"] {
        background: #111827;
    }
    [data-testid="stSidebar"] * {
        color: #e5e7eb;
    }
    .hero {
        padding: 1.6rem 1.8rem;
        border-radius: 22px;
        background: linear-gradient(135deg, #0b1f33 0%, #164a70 55%, #1c718c 100%);
        color: white;
        margin-bottom: 1.2rem;
    }
    .metric {
        background: #172033 !important;
        border-radius: 17px;
        padding: 1rem 1.1rem;
        border: 1px solid #263449;
        margin-bottom: 10px;
    }
    .metric-label { color: #aab7c7 !important; }
    .metric-value { color: #f8fafc !important; font-size: 1.7rem; font-weight: 700; }
    .section-title { color: #f1f5f9 !important; }
    .side-note { color: #aab7c7 !important; }
    div[data-testid="stChatMessage"] { border-radius: 16px; }
    div[data-testid="stChatInput"] {
        background: #171b26 !important;
        border: 1px solid #2b3445 !important;
    }
    div[data-testid="stChatInput"] textarea {
        background: #171b26 !important;
        color: #f1f5f9 !important;
    }
    .stButton > button {
        border-radius: 12px !important;
        min-height: 48px !important;
        font-weight: 650 !important;
        background: #172033 !important;
        color: #f1f5f9 !important;
        border: 1px solid #314057 !important;
    }
    .stButton > button:hover {
        background: #22304a !important;
        border-color: #4a6385 !important;
    }
</style>
""", unsafe_allow_html=True)

if "kb" not in st.session_state:
    st.session_state.kb = load_default()
if "messages" not in st.session_state:
    st.session_state.messages = []

kb = st.session_state.kb

with st.sidebar:
    st.markdown("## ⚓ Agente Tiara")
    st.caption("Centro de control del barco")
    st.divider()

    if kb:
        st.success("● Base activa")
        st.caption(f"Fuente: {kb['metadata']['source_file']}")
    else:
        st.warning("○ Sin base cargada")

    uploaded = st.file_uploader(
        "Actualizar fuente de datos",
        type=["xlsx", "xlsm"],
        help="Carga una nueva versión del Excel para reconstruir la base de conocimiento.",
    )
    if uploaded and st.button("Importar Excel", use_container_width=True):
        with st.spinner("Procesando Excel..."):
            st.session_state.kb = import_excel(uploaded)
            st.session_state.messages = []
        st.rerun()

    st.divider()
    st.markdown("**Resumen**")
    sidebar_records = sum(len(kb[g]) for g in (
        "maintenance_records", "log_records", "inventory_items", "fuel_records",
        "budget_records", "invoice_records", "cleaning_records", "documents"
    )) if kb else 0

    st.markdown(
        f'<div class="metric"><div class="metric-label">Hojas de datos</div>'
        f'<div class="metric-value">{len(kb["sheets"]) if kb else 0}</div></div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        f'<div class="metric"><div class="metric-label">Registros</div>'
        f'<div class="metric-value">{sidebar_records:,}</div></div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        f'<div class="metric"><div class="metric-label">Mantenimientos</div>'
        f'<div class="metric-value">{len(kb["maintenance_records"]) if kb else 0}</div></div>',
        unsafe_allow_html=True,
    )

    st.divider()
    st.markdown("**Módulos**")
    st.caption("🔧 Mantenimiento")
    st.caption("📋 Bitácora y eventos")
    st.caption("📦 Inventario")
    st.caption("⛽ Combustible")
    st.caption("🧾 Facturas y costos")
    st.caption("📄 Permisos y seguros")

    st.divider()
    st.markdown(
        '<div class="side-note">La información de operación proviene del Excel cargado como fuente de verdad.</div>',
        unsafe_allow_html=True,
    )

st.markdown("""
<div class="hero">
    <h1>⚓ AGENTE TIARA</h1>
    <p>Tu centro de información para mantenimiento, operación, inventario y documentación del barco.</p>
</div>
""", unsafe_allow_html=True)

if not kb:
    st.info("Carga el Excel de septiembre desde la barra lateral para activar el centro de control.")
    st.stop()

# El área principal queda libre para los accesos rápidos.
st.markdown('<div class="section-title">Asistente</div>', unsafe_allow_html=True)
st.markdown('<div class="section-title">Accesos rápidos</div>', unsafe_allow_html=True)

b1, b2, b3, b4, b5 = st.columns(5)
quick_query = None

with b1:
    if st.button("🛥️ Estado general", use_container_width=True):
        quick_query = "Dame el estado general del barco"

with b2:
    if st.button("🔧 Motores", use_container_width=True):
        quick_query = "¿Cuál fue el último mantenimiento de motores?"

with b3:
    if st.button("🔩 Kit 300 h", use_container_width=True):
        quick_query = "¿Cuál fue el último mantenimiento del kit de 300 horas?"

with b4:
    if st.button("⚙️ Patas", use_container_width=True):
        quick_query = "¿Cuál fue el último mantenimiento de patas de motores?"

with b5:
    if st.button("🔌 Generador", use_container_width=True):
        quick_query = "¿Cuál fue el último mantenimiento del generador?"


# Una sola interacción visible: cada pregunta nueva reemplaza la anterior.
prompt = st.chat_input("Escribe aquí tu consulta sobre el Tiara…")
if prompt:
    quick_query = prompt

if quick_query:
    response = answer(kb, quick_query)
    st.session_state.messages = [
        {"role": "user", "content": quick_query},
        {"role": "assistant", "content": response},
    ]

if st.session_state.messages:
    with st.chat_message("user"):
        st.markdown(st.session_state.messages[0]["content"])
    with st.chat_message("assistant"):
        st.markdown(st.session_state.messages[1]["content"])

st.markdown('<div class="section-title">Estado del barco</div>', unsafe_allow_html=True)

st.markdown('<div class="section-title">Fuente de datos</div>', unsafe_allow_html=True)

with st.container():
    st.markdown("#### 📚 Fuente de datos")
    st.write(f"**Archivo:** {kb['metadata']['source_file']}")
    st.write(f"**Actualizado:** {kb['metadata'].get('updated_at', 'No disponible')}")
    st.write(f"**Hojas:** {len(kb['sheets'])}")
    st.write(f"**Registros:** {sidebar_records:,}")

    with st.expander("Ver hojas"):
        for sheet_name in kb["sheets"]:
            st.write(f"• {sheet_name}")
