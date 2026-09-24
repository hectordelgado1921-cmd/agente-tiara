from __future__ import annotations

import json
import re
import unicodedata
from datetime import datetime, date
from zoneinfo import ZoneInfo
from pathlib import Path
import hashlib
import inspect
import tempfile
import zipfile
import base64
import urllib.request
import urllib.error

import streamlit as st

# ============================================================
# AGENTE TIARA 2.0
# Cerebro determinista basado en el Excel de septiembre 2026
# ============================================================

ROOT = Path(__file__).resolve().parent
APP_VERSION = "2.9.0"
APP_UPDATED = "14/09/2026"  # Fecha del código; la carga del Excel es dinámica.
APP_TIMEZONE = ZoneInfo("America/Costa_Rica")
DATA_DIR = ROOT / "data"
JSON_FILE = DATA_DIR / "base_conocimiento_tiara_septiembre_2026_v2.json"
JSON_FILE_OLD = DATA_DIR / "base_conocimiento_tiara_septiembre_2026.json"
EXCEL_CANDIDATES = [
    ROOT / "VITACORA DE BARCO  SEPTIEMBRE 2026.xlsx",
    ROOT / "bitácora barco septiembre 2026.xlsx",
    DATA_DIR / "VITACORA DE BARCO  SEPTIEMBRE 2026.xlsx",
    DATA_DIR / "bitácora barco septiembre 2026.xlsx",
]

MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
    "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
    "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}


# ============================================================
# VISIÓN / FOTOGRAFÍAS DEL EXCEL
# ============================================================
# La clave de Gemini se lee exclusivamente desde Streamlit Secrets.
# Nunca se muestra ni se guarda dentro de la base de conocimiento.

def _gemini_api_key():
    try:
        cfg = st.secrets.get("gemini", {})
        key = str(cfg.get("api_key", "")).strip()
        if key:
            return key
        # Compatibilidad opcional con un secret plano.
        key = str(st.secrets.get("GEMINI_API_KEY", "")).strip()
        return key or None
    except Exception:
        return None


def _prepare_image_for_vision(raw: bytes, fmt: str):
    """Reduce una foto si es demasiado grande, sin alterar la original del Excel."""
    if not raw:
        return raw, ("jpeg" if fmt == "jpg" else fmt)
    try:
        from PIL import Image, ImageOps
        import io
        img = Image.open(io.BytesIO(raw))
        img.load()
        try:
            img = ImageOps.exif_transpose(img)
        except Exception:
            pass
        # Mayor resolución para documentos con letra pequeña.
        max_side = 2200
        if max(img.size) > max_side:
            ratio = max_side / float(max(img.size))
            img = img.resize((max(1, int(img.width*ratio)), max(1, int(img.height*ratio))), Image.LANCZOS)
        out = io.BytesIO()
        if fmt.lower() in ("png",) and img.mode in ("RGBA", "LA"):
            img.save(out, format="PNG", optimize=True)
            return out.getvalue(), "png"
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        img.save(out, format="JPEG", quality=82, optimize=True)
        return out.getvalue(), "jpeg"
    except Exception:
        return raw, ("jpeg" if fmt == "jpg" else fmt)


# ============================================================
# COMPATIBILIDAD AUTOMÁTICA DE SERVICIOS GEMINI
# ============================================================
# Agente Tiara usa REST directamente: no depende de una versión instalada
# del SDK de Google. Los modelos se aíslan aquí para poder sustituirlos sin
# tocar el resto del agente.
GEMINI_VISION_MODEL = "gemini-3.5-flash-lite"
GEMINI_TRANSCRIBE_MODEL = "gemini-3.5-transcribe"
GEMINI_VISION_FALLBACKS = [
    "gemini-3.5-flash-lite",
    "gemini-3.6-flash",
    "gemini-3.5-flash",
]
GEMINI_TRANSCRIBE_FALLBACKS = [
    "gemini-3.5-transcribe",
]


def _gemini_models_url():
    return "https://generativelanguage.googleapis.com/v1beta/models"


def _gemini_list_models():
    """Obtiene modelos disponibles para esta API key.

    Es solo una capa de compatibilidad: si Google agrega un modelo nuevo,
    podemos detectarlo sin tener que instalar una librería nueva.
    """
    key = _gemini_api_key()
    if not key:
        return []
    req = urllib.request.Request(
        _gemini_models_url(),
        headers={"x-goog-api-key": key},
        method="GET",
    )
    try:
        with urllib.request.urlopen(req, timeout=20) as response:
            obj = json.loads(response.read().decode("utf-8"))
        return obj.get("models") or []
    except Exception:
        return []


def _gemini_compatible_model(kind: str, preferred: str, fallbacks):
    """Devuelve primero el modelo actual; solo busca otro si no está disponible.

    Así una nueva versión de Google no cambia silenciosamente el comportamiento
    mientras el modelo actual siga funcionando. Si deja de estar disponible,
    se intenta un reemplazo compatible y, finalmente, un modelo descubierto
    dinámicamente en la lista de modelos de Google.
    """
    models = _gemini_list_models()
    available = {
        str(m.get("name", "")).split("/models/")[-1]: m
        for m in models
        if m.get("name")
    }

    candidates = []
    for model in [preferred] + list(fallbacks):
        if model not in candidates:
            candidates.append(model)
        if model in available:
            methods = available[model].get("supportedGenerationMethods") or []
            if "generateContent" in methods:
                return model

    # Descubrimiento futuro: solo modelos estables/compatibles de la misma
    # familia funcional, evitando preview/experimental y modelos de imagen.
    discovered = []
    for name, meta in available.items():
        low = name.lower()
        methods = meta.get("supportedGenerationMethods") or []
        if "generateContent" not in methods:
            continue
        if "preview" in low or "experimental" in low or "image" in low:
            continue
        if kind == "vision" and not ("flash-lite" in low or "flash" in low):
            continue
        if kind == "transcribe" and "transcribe" not in low:
            continue
        discovered.append(name)

    # Orden descendente por nombre para preferir una generación nueva dentro
    # de la familia compatible. Solo se llega aquí cuando los candidatos
    # conocidos ya no están disponibles.
    discovered.sort(reverse=True)
    return discovered[0] if discovered else None


def _gemini_generate_content(model: str, payload: dict, timeout: int):
    key = _gemini_api_key()
    if not key:
        return None, "NO_API_KEY"
    endpoint = f"{_gemini_models_url()}/{model}:generateContent"
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        endpoint,
        data=body,
        headers={
            "Content-Type": "application/json",
            "x-goog-api-key": key,
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as response:
            result = json.loads(response.read().decode("utf-8"))
        parts = (((result.get("candidates") or [{}])[0].get("content") or {}).get("parts") or [])
        text = " ".join(
            str(x.get("text", "")).strip() for x in parts if x.get("text")
        ).strip()
        return text or None, None
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        return None, f"HTTP {exc.code}: {detail[:500]}"
    except Exception as exc:
        return None, str(exc)


def _gemini_vision_request(image_bytes: bytes, mime_type: str, prompt: str):
    key = _gemini_api_key()
    if not key:
        return None, "NO_API_KEY"
    payload = {
        "contents": [{"parts": [
            {"text": prompt},
            {"inline_data": {
                "mime_type": mime_type,
                "data": base64.b64encode(image_bytes).decode("ascii"),
            }},
        ]}],
        "generationConfig": {"temperature": 0.1, "maxOutputTokens": 900},
    }
    preferred = GEMINI_VISION_MODEL
    # Primero usamos el modelo conocido. Si Google lo retira, descubrimos un
    # reemplazo compatible sin modificar el resto de la aplicación.
    model = _gemini_compatible_model("vision", preferred, GEMINI_VISION_FALLBACKS)
    if not model:
        return None, "NO_COMPATIBLE_VISION_MODEL"
    text, err = _gemini_generate_content(model, payload, 45)
    if text:
        return text, None
    # Si el catálogo estaba desactualizado, prueba los candidatos conocidos
    # una vez antes de devolver el error.
    for candidate in GEMINI_VISION_FALLBACKS:
        if candidate == model:
            continue
        text, candidate_err = _gemini_generate_content(candidate, payload, 45)
        if text:
            return text, None
        err = candidate_err or err
    return None, err


def _gemini_upload_file(audio_bytes: bytes, mime_type: str):
    """Sube el audio a Gemini Files API y devuelve su URI."""
    key = _gemini_api_key()
    if not key:
        return None, "NO_API_KEY"
    if not audio_bytes:
        return None, "NO_AUDIO"

    mime = (mime_type or "audio/webm").split(";")[0].strip().lower() or "audio/webm"
    size = len(audio_bytes)
    base = "https://generativelanguage.googleapis.com"

    start_url = f"{base}/upload/v1beta/files"
    start_payload = json.dumps({
        "file": {"display_name": "tiara_voice_question"}
    }).encode("utf-8")
    start_req = urllib.request.Request(
        start_url,
        data=start_payload,
        headers={
            "x-goog-api-key": key,
            "X-Goog-Upload-Protocol": "resumable",
            "X-Goog-Upload-Command": "start",
            "X-Goog-Upload-Header-Content-Length": str(size),
            "X-Goog-Upload-Header-Content-Type": mime,
            "Content-Type": "application/json",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(start_req, timeout=30) as response:
            upload_url = response.headers.get("X-Goog-Upload-URL") or response.headers.get("x-goog-upload-url")
        if not upload_url:
            return None, "NO_UPLOAD_URL"

        upload_req = urllib.request.Request(
            upload_url,
            data=audio_bytes,
            headers={
                "Content-Length": str(size),
                "X-Goog-Upload-Offset": "0",
                "X-Goog-Upload-Command": "upload, finalize",
            },
            method="POST",
        )
        with urllib.request.urlopen(upload_req, timeout=90) as response:
            result = json.loads(response.read().decode("utf-8"))

        file_obj = result.get("file") or {}
        uri = file_obj.get("uri")
        if not uri:
            return None, f"NO_FILE_URI: {str(result)[:400]}"
        return uri, None
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        return None, f"UPLOAD_HTTP_{exc.code}: {detail[:500]}"
    except Exception as exc:
        return None, f"UPLOAD_ERROR: {exc}"


def _gemini_interaction_transcribe(file_uri: str, mime_type: str):
    """Transcribe con el endpoint oficial Interactions de Gemini 3.5 Transcribe."""
    key = _gemini_api_key()
    if not key:
        return None, "NO_API_KEY"

    mime = (mime_type or "audio/webm").split(";")[0].strip().lower() or "audio/webm"
    payload = {
        "model": GEMINI_TRANSCRIBE_MODEL,
        "input": [{
            "type": "audio",
            "uri": file_uri,
            "mime_type": mime,
        }],
        "generation_config": {
            "transcription_config": {
                "language_codes": ["es-CR"],
                "custom_vocabulary": [
                    "Tiara", "Mercury", "Cummins", "Onan", "Racor",
                    "horómetro", "horómetros", "motores", "patas",
                    "kit 300 horas", "300 horas", "Generador Cummins Onan"
                ],
            }
        },
    }
    req = urllib.request.Request(
        "https://generativelanguage.googleapis.com/v1beta/interactions",
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "x-goog-api-key": key,
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=90) as response:
            result = json.loads(response.read().decode("utf-8"))

        text = result.get("output_text")
        if isinstance(text, str) and text.strip():
            return text.strip(), None

        for item in result.get("outputs") or []:
            if item.get("type") == "text" and item.get("text"):
                return str(item["text"]).strip(), None

        for step in result.get("steps") or []:
            for part in step.get("content") or []:
                if part.get("type") == "text" and part.get("text"):
                    return str(part["text"]).strip(), None

        return None, f"NO_TRANSCRIPTION_TEXT: {str(result)[:500]}"
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        return None, f"INTERACTION_HTTP_{exc.code}: {detail[:500]}"
    except Exception as exc:
        return None, f"INTERACTION_ERROR: {exc}"


def _gemini_audio_transcribe(audio_bytes: bytes, mime_type: str = "audio/webm"):
    """Transcribe una pregunta hablada usando Gemini.

    Ruta principal: Files API -> Interactions -> Gemini 3.5 Transcribe.
    Ruta de respaldo: generateContent con audio inline usando Gemini Flash.
    Ambas devuelven SOLO el texto de la pregunta para pasarlo al mismo run_agent().
    """
    key = _gemini_api_key()
    if not key:
        return None, "NO_API_KEY"
    if not audio_bytes:
        return None, "NO_AUDIO"
    if len(audio_bytes) > 20 * 1024 * 1024:
        return None, "AUDIO_TOO_LARGE"

    mime = (mime_type or "audio/webm").split(";")[0].strip().lower() or "audio/webm"

    # 1) Ruta oficial dedicada para voz a texto.
    file_uri, upload_error = _gemini_upload_file(audio_bytes, mime)
    if file_uri:
        text, err = _gemini_interaction_transcribe(file_uri, mime)
        if text:
            return text, None
    else:
        err = upload_error

    # 2) Respaldo: comprensión de audio mediante generateContent.
    # Esto evita que una diferencia temporal de disponibilidad del endpoint
    # Interactions deje sin funcionamiento el micrófono.
    fallback_models = ["gemini-3.6-flash", "gemini-3.5-flash"]
    payload = {
        "contents": [{"parts": [
            {"text": (
                "Transcribe exactamente la pregunta hablada en español. "
                "Devuelve únicamente el texto de la pregunta, sin explicaciones, "
                "sin comillas y sin responderla. Conserva estos términos técnicos: "
                "Tiara, Mercury, Cummins, Onan, Racor, horómetro, motores, patas, "
                "kit 300 horas y generador."
            )},
            {"inline_data": {
                "mime_type": mime,
                "data": base64.b64encode(audio_bytes).decode("ascii"),
            }},
        ]}],
        "generationConfig": {"temperature": 0.0, "maxOutputTokens": 300},
    }

    fallback_err = err
    for model in fallback_models:
        text, candidate_err = _gemini_generate_content(model, payload, 90)
        if text:
            return text, None
        fallback_err = candidate_err or fallback_err

    return None, fallback_err or "NO_TRANSCRIPTION"

def _analyze_embedded_images(images):
    """Interpreta fotografías incrustadas y devuelve solo texto estructurado.

    Las imágenes binarias se eliminan después del análisis para no inflar el
    snapshot de Supabase. La fotografía original sigue siendo la fuente en el
    Excel del usuario.
    """
    if not images:
        return images, {"enabled": False, "analyzed": 0, "reason": "sin_fotos"}
    if not _gemini_api_key():
        # No enviamos ninguna imagen si el usuario todavía no configuró la API.
        clean = []
        for im in images:
            clean.append({k:v for k,v in im.items() if k != "data_base64"})
        return clean, {"enabled": False, "analyzed": 0, "reason": "sin_api_key"}

    # Un Excel del barco puede contener cientos de imágenes decorativas o
    # adjuntas en otras hojas. Analizar todas con Gemini durante la carga
    # bloquea Streamlit durante muchos minutos. La visión de documentos que
    # necesita el agente está concentrada en PERMISOS & SEGUROS.
    targets = [im for im in images if str(im.get("sheet", "")) == "PERMISOS & SEGUROS"]
    if not targets:
        targets = images[:10]
    max_images = 10
    targets = targets[:max_images]

    prompt = (
        "Eres el módulo de visión del Agente Tiara. Analiza esta fotografía "
        "incrustada en un Excel de permisos/seguros de una embarcación. "
        "Extrae SOLO información realmente visible: tipo de documento, "
        "nombre de embarcación/empresa/persona si se lee, número de póliza, "
        "permiso/certificado si se lee, fechas de emisión y vencimiento, "
        "aseguradora/autoridad, y cualquier dato importante visible. "
        "Si algo no se puede leer, dilo como 'no legible'. No inventes ni "
        "completes datos faltantes. Responde en español y en formato breve."
    )

    def analyze_one(im):
        item = {k:v for k,v in im.items() if k != "data_base64"}
        b64 = im.get("data_base64")
        if not b64:
            item["vision_analysis"] = "Sin imagen binaria disponible."
            item["vision_error"] = "NO_IMAGE_DATA"
            return item
        try:
            raw = base64.b64decode(b64)
            fmt = str(im.get("format") or "jpeg").lower()
            raw2, fmt2 = _prepare_image_for_vision(raw, fmt)
            mime = "image/png" if fmt2 == "png" else ("image/webp" if fmt2 == "webp" else "image/jpeg")
            text, err = _gemini_vision_request(raw2, mime, prompt)
            item["vision_analysis"] = text or "No se pudo interpretar visualmente esta fotografía."
            item["vision_error"] = err if err and err != "NO_API_KEY" else None
        except Exception as exc:
            item["vision_analysis"] = "No se pudo procesar esta fotografía."
            item["vision_error"] = str(exc)
        return item

    # Cuatro solicitudes simultáneas reducen drásticamente el tiempo total
    # sin disparar una cantidad excesiva de llamadas a la API.
    from concurrent.futures import ThreadPoolExecutor, as_completed
    analyzed_map = {}
    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = {executor.submit(analyze_one, im): idx for idx, im in enumerate(targets)}
        for future in as_completed(futures):
            analyzed_map[futures[future]] = future.result()

    analyzed = [analyzed_map[i] for i in range(len(targets))]

    # Conserva metadata de las demás fotos, sin bytes ni análisis, para no
    # inflar la base de Supabase ni bloquear la carga del Excel.
    target_ids = {id(im) for im in targets}
    for im in images:
        if id(im) not in target_ids:
            analyzed.append({k:v for k,v in im.items() if k != "data_base64"})

    return analyzed, {"enabled": True, "analyzed": len(targets), "total": len(images),
                     "vision_scope": "PERMISOS & SEGUROS"}


# ============================================================
# PERSISTENCIA SUPABASE
# ============================================================
# La aplicación corre en el servidor de Streamlit. La clave de Supabase
# se lee exclusivamente desde st.secrets y nunca se muestra en pantalla.
SUPABASE_TABLE = "tiara_data"

def _supabase_credentials():
    """Lee la configuración de Supabase sin exponer secretos.

    La aplicación funciona con la clave pública que el usuario ya configuró
    como `supabase.key`. Si en el futuro existe `service_key`, se acepta como
    respaldo, pero nunca se exige para arrancar ni para procesar el Excel.
    """
    try:
        cfg = st.secrets.get("supabase", {})
        url = str(cfg.get("url", "")).strip().rstrip("/")
        key = str(cfg.get("key", "")).strip()
        if not key:
            key = str(cfg.get("service_key", "")).strip()
        if url and key:
            return url, key
    except Exception:
        pass
    return None, None


def _supabase_request(method, path, payload=None, query=""):
    import urllib.request
    import urllib.error

    base_url, key = _supabase_credentials()
    if not base_url or not key:
        raise RuntimeError("No están configurados los Secrets de Supabase en Streamlit.")

    url = f"{base_url}/rest/v1/{SUPABASE_TABLE}{path}{query}"
    headers = {
        "apikey": key,
        "Content-Type": "application/json",
        "Prefer": "return=minimal",
    }
    body = None if payload is None else json.dumps(payload, ensure_ascii=False, allow_nan=False).encode("utf-8")
    request = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            raw = response.read().decode("utf-8")
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase respondió HTTP {exc.code}: {detail[:500]}") from exc
    except Exception as exc:
        # Streamlit Cloud puede tardar en recibir la respuesta de Supabase
        # aunque el servidor sí esté procesando el snapshot. Reintentamos
        # una sola vez en operaciones de escritura para tolerar ese timeout.
        if method.upper() in {"POST", "PATCH", "PUT"} and "timed out" in str(exc).lower():
            try:
                with urllib.request.urlopen(request, timeout=180) as response:
                    raw = response.read().decode("utf-8")
                    return json.loads(raw) if raw else None
            except urllib.error.HTTPError as retry_exc:
                detail = retry_exc.read().decode("utf-8", errors="replace")
                raise RuntimeError(f"Supabase respondió HTTP {retry_exc.code}: {detail[:500]}") from retry_exc
            except Exception as retry_exc:
                raise RuntimeError(f"No se pudo conectar con Supabase tras reintentar: {retry_exc}") from retry_exc
        raise RuntimeError(f"No se pudo conectar con Supabase: {exc}") from exc


def load_supabase_data():
    """Carga el último snapshot activo guardado en Supabase."""
    try:
        rows = _supabase_request(
            "GET",
            "",
            query=(
                "?select=data,source_filename,source_hash,version_number,created_at"
                "&is_active=eq.true&order=created_at.desc&limit=1"
            ),
        )
        if not rows:
            return None
        row = rows[0]
        data = row.get("data")
        if not isinstance(data, dict):
            return None
        # Supabase puede contener un snapshot antiguo (base 1.0).
        # Normalizarlo aquí es imprescindible para que el cerebro 2.0
        # encuentre mantenimiento, horómetros, combustible, presupuesto,
        # inventario y demás registros con la estructura actual.
        data = normalize_loaded_data(data)
        meta = data.setdefault("metadata", {})
        if row.get("source_filename") and not meta.get("source_file"):
            meta["source_file"] = row["source_filename"]
        if row.get("source_hash") and not meta.get("upload_sha256"):
            meta["upload_sha256"] = row["source_hash"]
        if row.get("version_number") is not None and not meta.get("version"):
            # version_number es un entero en Supabase; conservamos la versión
            # legible del snapshot cuando ya viene dentro de metadata.
            meta["version_number"] = row["version_number"]
        meta["supabase_persisted"] = True
        meta["supabase_created_at"] = row.get("created_at")
        return data
    except Exception as exc:
        st.session_state.supabase_load_error = str(exc)
        return None


def save_supabase_data(data):
    """Guarda un snapshot completo de la base normalizada del Excel."""
    meta = data.get("metadata", {})
    payload = {
        "source_filename": meta.get("source_file", "Excel Tiara"),
        "source_hash": meta.get("upload_sha256"),
        "version_number": 210,
        "is_active": True,
        "data": data,
    }
    _supabase_request("POST", "", payload=payload)


MONTH_NAMES = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
    5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
    9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre",
}

STOPWORDS = {
    "el", "la", "los", "las", "de", "del", "un", "una", "unos", "unas",
    "que", "qué", "se", "ha", "han", "le", "les", "al", "en", "por",
    "para", "con", "y", "o", "a", "lo", "me", "nos", "su", "sus",
    "fue", "ser", "es", "son", "hay", "cuál", "cual", "cuáles", "cuales",
    "cuando", "cuándo", "como", "cómo", "toca", "tiene", "tienen", "dame",
    "quiero", "saber", "sobre", "del", "durante", "este", "esta", "del",
}


def norm(value: object) -> str:
    if value is None:
        return ""
    s = str(value).strip().lower()
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"\s+", " ", s)
    return s


def safe_float(value: object):
    if value is None or value == "":
        return None
    try:
        x = float(value)
        return x if x == x else None
    except Exception:
        s = str(value).replace("$", "").replace(",", "").strip()
        try:
            return float(s)
        except Exception:
            return None


def parse_date(value: object):
    if isinstance(value, (datetime, date)):
        return value.date() if isinstance(value, datetime) else value
    if not value:
        return None
    s = str(value).strip()

    # Los valores de fecha del Excel se serializan con ``datetime.isoformat()``,
    # por lo que pueden llegar como ``2026-08-12T00:00:00``. Antes solo se
    # aceptaba ``2026-08-12`` y esas fechas terminaban como ``date.min``. Eso
    # hacía que el desempate por fila colocara un registro histórico (1075 h)
    # por encima de lecturas recientes (2234 h).
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).date()
    except Exception:
        pass

    for fmt in ("%d-%m-%y", "%d-%m-%Y", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            pass
    return None


def extract_dates(value: object):
    out = []
    if value is None:
        return out
    for part in re.split(r"[|\n]+", str(value)):
        d = parse_date(part.strip())
        if d:
            out.append(d)
    return out


def extract_hours(value: object):
    out = []
    if value is None:
        return out
    for part in re.split(r"[|\n]+", str(value)):
        part = part.strip()
        if re.fullmatch(r"\d+(?:\.0)?", part):
            try:
                out.append(int(float(part)))
            except Exception:
                pass
    return out


def fmt_date(d):
    return d.strftime("%d-%m-%Y") if d else "NO DETERMINADO"


def money(v):
    return f"${v:,.2f} USD"


@st.cache_data(show_spinner=False)
def load_data():
    # 1) Preferir la base v2 si existe.
    if JSON_FILE.exists():
        with JSON_FILE.open("r", encoding="utf-8") as f:
            return normalize_loaded_data(json.load(f))

    # 2) Si está disponible el Excel original, reconstruir la base desde la fuente.
    for path in EXCEL_CANDIDATES:
        if path.exists():
            return build_from_excel(path)

    # 3) Compatibilidad con la base JSON antigua que ya está en GitHub.
    if JSON_FILE_OLD.exists():
        with JSON_FILE_OLD.open("r", encoding="utf-8") as f:
            return normalize_loaded_data(json.load(f))

    return None


def excel_serial_to_date(value):
    if value is None:
        return None
    try:
        n = float(str(value).strip())
    except Exception:
        return None
    # Fechas seriales de Excel (sistema 1900).
    if 20000 <= n <= 80000:
        from datetime import timedelta
        return date(1899, 12, 30) + timedelta(days=int(n))
    return None


def normalize_loaded_data(data):
    """Convierte la base antigua 1.0 a la estructura que usa el agente 2.0."""
    if not data:
        return data

    # Ya está normalizada.
    if "inventory_records" in data and "documents_records" in data:
        return data

    out = {
        "metadata": dict(data.get("metadata", {})),
        "sheets": dict(data.get("sheets", {})),
        "maintenance_records": [],
        "log_records": [],
        "inventory_records": [],
        "fuel_records": [],
        "budget_records": [],
        "invoice_records": [],
        "cleaning_records": [],
        "documents_records": [],
        "images": [],
    }

    out["metadata"]["version"] = APP_VERSION

    # Mantenimiento: normalizar los nombres de columnas de la base 1.0.
    for rec in data.get("maintenance_records", []):
        d = rec.get("data", {})
        def first_key(*keys):
            for k in keys:
                if k in d and d[k] not in (None, ""):
                    return d[k]
            return None

        dates_raw = first_key("fecha", "fecha cambio o mantenimiento", "FECHA")
        dates = []
        for part in re.split(r"[|\n]+", str(dates_raw or "")):
            part = part.strip()
            dt = parse_date(part) or excel_serial_to_date(part)
            if dt:
                dates.append(dt.isoformat())

        hours_raw = first_key("historial_horas", "historial mantenimientos ", "historial mantenimientos", "HISTORIAL MANTENIMIENTOS")
        interval = first_key("cada_horas", "cada (horas)", "cada horas", "CADA (HORAS)")
        normalized = {
            "mantenimiento": first_key("mantenimiento", "MANTENIMIENTO"),
            "cada_horas": interval,
            "historial_horas": hours_raw,
            "fecha": "|".join(dates),
            "observaciones": first_key("observaciones", "OBSERVACIONES"),
        }
        out["maintenance_records"].append({
            "source_sheet": rec.get("source_sheet", "CHECKLIST MANT. PREVENTIVO"),
            "source_row": rec.get("source_row"),
            "data": normalized,
        })

    # Los demás módulos conservan sus datos y además reciben un campo values
    # para que el cerebro 2.0 pueda consultarlos sin romper la base antigua.
    mappings = {
        "log_records": "log_records",
        "inventory_records": "inventory_items",
        "fuel_records": "fuel_records",
        "budget_records": "budget_records",
        "invoice_records": "invoice_records",
        "cleaning_records": "cleaning_records",
        "documents_records": "documents",
    }
    for target, old_key in mappings.items():
        for rec in data.get(old_key, []):
            vals = list((rec.get("data") or {}).values())
            if target == "log_records":
                # Reinterpretar la bitácora antigua con la nueva estructura
                # para que el estado, sistema y horas queden disponibles.
                parsed = parse_log_record(vals, rec.get("source_row"))
                parsed["source_sheet"] = rec.get("source_sheet")
                out[target].append(parsed)
            else:
                out[target].append({
                    "source_sheet": rec.get("source_sheet"),
                    "source_row": rec.get("source_row"),
                    "values": vals,
                    "data": rec.get("data", {}),
                })

    return out


def _find_sheet(wb, expected_name):
    """Encuentra una hoja tolerando tildes, espacios y pequeñas variaciones."""
    target = norm(expected_name).replace(" ", "")
    for ws in wb.worksheets:
        candidate = norm(ws.title).replace(" ", "")
        if candidate == target:
            return ws
    # Coincidencia parcial como respaldo para nombres que cambian ligeramente.
    for ws in wb.worksheets:
        candidate = norm(ws.title).replace(" ", "")
        if target in candidate or candidate in target:
            return ws
    return None


def _read_sheet_rows(ws):
    """Lee una hoja en modo read_only y devuelve solo los valores serializados."""
    if ws is None:
        return []
    rows = []
    for row in ws.iter_rows(values_only=True):
        vals = [serialize(v) for v in row]
        if any(v not in (None, "") for v in vals):
            rows.append(vals)
    return rows



# ============================================================
# BITÁCORA ESTRUCTURADA
# ============================================================

LOG_HEADERS = {
    "observation": 0,
    "report_date": 1,
    "pending": 2,
    "done": 3,
    "installation_date": 4,
    "hours": 5,
    "notes": 6,
}


def _is_checked(value):
    """Reconoce casillas/checks que puedan llegar como texto o booleano."""
    if value is True:
        return True
    if value in (1, 1.0):
        return True
    n = norm(value)
    return n in {"true", "verdadero", "si", "sí", "yes", "checked", "check", "✓", "✔", "☑"}


def _is_x(value):
    if value is None:
        return False
    return norm(value) in {"x", "✕", "✖", "×"}


def _log_hours_from_value(value):
    """Acepta 540, '540 h', etc., pero ignora números que parezcan fechas."""
    if value in (None, ""):
        return None
    try:
        x = float(value)
        if 1 <= x <= 10000:
            return int(x) if x.is_integer() else x
    except Exception:
        pass
    text = str(value).replace(",", ".")
    normalized = norm(text)
    # "kit 300 h" es un intervalo de servicio, no un horómetro actual.
    if re.search(r"\bkit\s+(?:de\s+)?300\s*(?:h|hr|hrs|hora|horas)?\b", normalized):
        return None
    m = re.search(r"(?<!\d)(\d+(?:\.\d+)?)\s*(?:h|hr|hrs|hora|horas)\b", normalized)
    if m:
        x = float(m.group(1))
        return int(x) if x.is_integer() else x
    return None


def _log_text_hours(text):
    """Extrae horas explícitas escritas dentro de observaciones/notas."""
    text = norm(text)
    matches = re.finditer(r"(?<!\d)(\d+(?:\.\d+)?)\s*(?:h|hr|hrs|hora|horas)\b", text)
    out = []
    for match in matches:
        # No tomar "kit 300 h" como horómetro.
        prefix = text[max(0, match.start()-12):match.start()]
        if re.search(r"\bkit\s+(?:de\s+)?$", prefix):
            continue
        x = float(match.group(1))
        if 1 <= x <= 10000:
            out.append(int(x) if x.is_integer() else x)
    # También reconoce el formato '(530) h de motor'.
    for m in re.findall(r"\((\d+(?:\.\d+)?)\)\s*h", text):
        x = float(m)
        if 1 <= x <= 10000:
            out.append(int(x) if x.is_integer() else x)
    return out


def _classify_log_system(text):
    """Clasifica la fila sin mezclar horas de motor con horas de generador."""
    p = norm(text)
    # Usamos palabras completas: "funcionando" no puede activar "onan" y
    # otras coincidencias accidentales no deben cambiar el sistema.
    has_generator = bool(re.search(r"\b(?:generador|cummins|onan)\b", p))
    has_motor = bool(re.search(r"\b(?:motor|motores|mercury|estribor|babor)\b", p))
    has_leg = bool(re.search(r"\b(?:pata|patas|outdrive|muelas)\b", p))

    # Si una misma fila habla explícitamente de motor y generador y solo hay
    # una cifra de horas, no adivinamos: esa hora NO se asigna a ninguno.
    if has_generator and has_motor:
        return "Ambos"
    if has_generator:
        return "Generador Cummins Onan"
    if has_leg:
        return "Patas Mercury"
    if has_motor:
        return "Motores Mercury"
    return "General"


def _log_status(pending, done, installation_date, text):
    """Regla de estado de la bitácora.

    C = X significa pendiente. D marcado significa resuelto/listo.
    La fecha de instalación/reparación en E confirma que el evento fue
    atendido, incluso cuando Excel entrega las casillas como objetos visuales
    sin valor de celda para openpyxl.
    """
    text_n = norm(text)
    if _is_checked(done) or installation_date is not None:
        return "RESUELTO"
    if _is_x(pending):
        return "PENDIENTE"
    if any(x in text_n for x in ("pendiente", "falta", "a la espera", "esperando")):
        return "PENDIENTE"
    return "REGISTRADO"


def parse_log_record(vals, row_number):
    observation = vals[LOG_HEADERS["observation"]] if len(vals) > 0 else None
    report_date = vals[LOG_HEADERS["report_date"]] if len(vals) > 1 else None
    pending = vals[LOG_HEADERS["pending"]] if len(vals) > 2 else None
    done = vals[LOG_HEADERS["done"]] if len(vals) > 3 else None
    installation_date = vals[LOG_HEADERS["installation_date"]] if len(vals) > 4 else None
    hours_value = vals[LOG_HEADERS["hours"]] if len(vals) > 5 else None
    notes = vals[LOG_HEADERS["notes"]] if len(vals) > 6 else None

    combined = " ".join(str(x) for x in (observation, notes) if x not in (None, ""))
    hours = _log_hours_from_value(hours_value)
    text_hours = _log_text_hours(combined)
    if hours is None and text_hours:
        hours = text_hours[-1]

    # La columna F puede contener una hora o, si algún día se escriben varias
    # en la misma celda, valores identificados por sistema. Esto permite
    # mantener separados motores y generador sin adivinar.
    hours_por_sistema = {}
    hours_text = norm(hours_value)
    labeled = re.findall(
        r"(\d+(?:\.\d+)?)\s*(?:h|hr|hrs|hora|horas)?\s*(?:de\s+)?"
        r"(motor(?:es)?|mercury|generador|cummins|onan|pata(?:s)?)\b",
        hours_text,
    )
    for raw_h, label in labeled:
        x = float(raw_h)
        if not (1 <= x <= 10000):
            continue
        label_n = norm(label)
        if label_n in {"generador", "cummins", "onan"}:
            key = "Generador Cummins Onan"
        elif label_n in {"pata", "patas"}:
            key = "Patas Mercury"
        else:
            key = "Motores Mercury"
        hours_por_sistema[key] = int(x) if x.is_integer() else x

    # Si la hora no viene etiquetada, se asigna únicamente cuando el texto de
    # la fila identifica un solo sistema. Si menciona dos sistemas, no se
    # adivina y esa cifra no contamina ningún horómetro.
    system = _classify_log_system(combined)
    if hours is not None and not hours_por_sistema and system in {
        "Motores Mercury", "Patas Mercury", "Generador Cummins Onan"
    }:
        hours_por_sistema[system] = hours

    status = _log_status(pending, done, installation_date, combined)

    return {
        "source_row": row_number,
        "values": vals,
        "data": {
            "observacion": observation,
            "fecha_reporte": report_date,
            "pendiente": pending,
            "listo": done,
            "fecha_instalacion": installation_date,
            "horas": hours,
            "horas_por_sistema": hours_por_sistema,
            "observaciones": notes,
            "sistema": system,
            "estado": status,
        },
    }


def log_records(data):
    return [r for r in data.get("log_records", []) if r.get("data")]


def _log_date(rec, key="fecha_reporte"):
    d = (rec.get("data") or {}).get(key)
    return parse_date(d) or excel_serial_to_date(d)


def _horometer_system(system):
    """Devuelve el sistema real del horómetro.

    Patas y Kit de 300 horas utilizan el mismo horómetro de motores.
    El generador mantiene su horómetro independiente.
    """
    if system in {"Motores Mercury", "Patas Mercury"}:
        return "Motores Mercury"
    return system


def latest_log_hours(data, system):
    """Obtiene el horómetro actual desde la Bitácora del Barco en Línea.

    Fuente de verdad:
      1) La hoja cruda ``VITACORA BARCO EN LINEA`` del snapshot, cuando existe.
      2) ``log_records`` normalizados, como compatibilidad con bases antiguas.
      3) El Checklist se usa únicamente como respaldo desde ``current_hours``.

    La lectura actual se determina por el evento más reciente para el sistema,
    no por el número de horas más alto. Motores y patas comparten horómetro;
    el generador mantiene uno independiente.
    """
    target = _horometer_system(system)

    def candidates_from_raw_rows(rows):
        candidates = []
        for row_number, vals in enumerate(rows, 1):
            if not isinstance(vals, (list, tuple)) or not vals:
                continue
            vals = list(vals)
            if len(vals) < 7:
                vals += [None] * (7 - len(vals))
            parsed = parse_log_record(vals, row_number)
            d = parsed.get("data") or {}
            horas_por_sistema = d.get("horas_por_sistema") or {}

            if target in horas_por_sistema:
                h = horas_por_sistema[target]
            elif system == "Patas Mercury" and "Patas Mercury" in horas_por_sistema:
                h = horas_por_sistema["Patas Mercury"]
            elif d.get("sistema") == target:
                h = d.get("horas")
            else:
                continue

            if h is None:
                continue
            rd = _log_date(parsed) or date.min
            candidates.append((rd, row_number, float(h)))
        return candidates

    # PRIORIDAD ABSOLUTA: si el snapshot conserva la hoja original, volvemos
    # a interpretar sus filas con la lógica vigente. Esto evita que una
    # clasificación vieja almacenada en log_records haga reaparecer, por
    # ejemplo, el histórico 1075 h como si fuera el horómetro actual.
    sheets = data.get("sheets") or {}
    raw_rows = None
    for sheet_name, sheet_data in sheets.items():
        if norm(sheet_name).replace(" ", "") == norm("VITACORA BARCO EN LINEA").replace(" ", ""):
            if isinstance(sheet_data, dict):
                raw_rows = sheet_data.get("rows")
            break

    raw_candidates = candidates_from_raw_rows(raw_rows or [])
    if raw_candidates:
        raw_candidates.sort(key=lambda x: (x[0], x[1]))
        return raw_candidates[-1][2]

    # COMPATIBILIDAD: snapshots que no conservan la hoja cruda.
    candidates = []
    for rec in log_records(data):
        raw_values = rec.get("values")
        if isinstance(raw_values, (list, tuple)) and raw_values:
            parsed = parse_log_record(list(raw_values), rec.get("source_row"))
            d = parsed.get("data") or {}
            rd = _log_date(parsed) or date.min
            row = int(parsed.get("source_row") or rec.get("source_row") or 0)
        else:
            d = rec.get("data") or {}
            rd = _log_date(rec) or date.min
            row = int(rec.get("source_row") or 0)

        horas_por_sistema = d.get("horas_por_sistema") or {}
        if target in horas_por_sistema:
            h = horas_por_sistema[target]
        elif system == "Patas Mercury" and "Patas Mercury" in horas_por_sistema:
            h = horas_por_sistema["Patas Mercury"]
        elif d.get("sistema") == target:
            h = d.get("horas")
        else:
            continue

        if h is None:
            continue
        candidates.append((rd, row, float(h)))

    if not candidates:
        return None
    candidates.sort(key=lambda x: (x[0], x[1]))
    return candidates[-1][2]


def _format_log_record(rec):
    d = rec.get("data") or {}
    report = _log_date(rec)
    install = _log_date(rec, "fecha_instalacion")
    parts = []
    if report:
        parts.append(f"reporte {fmt_date(report)}")
    if d.get("estado"):
        parts.append(d["estado"])
    if install:
        parts.append(f"atendido/instalado {fmt_date(install)}")
    if d.get("horas") is not None:
        parts.append(f"{d['horas']:g} h")
    system = d.get("sistema")
    if system and system != "General":
        parts.append(system)
    text = str(d.get("observacion") or "").strip()
    notes = str(d.get("observaciones") or "").strip()
    if notes:
        text += f" — {notes}"
    return " — ".join([text] + parts) if text else " | ".join(parts)


def _extract_excel_images(path: Path):
    """Extrae fotografías incrustadas del Excel y las asocia a hoja/celda.

    Las imágenes no forman parte de los valores de las celdas. Se conservan
    como evidencia visual para que el cerebro pueda localizarlas y, cuando
    sea posible, aplicar OCR. No se inventa contenido visual.
    """
    import openpyxl
    images = []
    try:
        wb_img = openpyxl.load_workbook(path, data_only=False, read_only=False)
        for ws in wb_img.worksheets:
            # Solo extraemos los bytes de las fotos de documentos; las demás
            # hojas pueden contener cientos de imágenes que no necesitan
            # visión durante la carga.
            is_document_sheet = ws.title == "PERMISOS & SEGUROS"
            for idx, img in enumerate(getattr(ws, "_images", []) or [], 1):
                try:
                    raw = img._data() if is_document_sheet else b""
                    fmt = str(getattr(img, "format", None) or "png").lower()
                    if fmt == "jpg": fmt = "jpeg"
                    anchor = getattr(img, "anchor", None)
                    marker = getattr(anchor, "_from", None)
                    row = (getattr(marker, "row", 0) + 1) if marker is not None else None
                    col = (getattr(marker, "col", 0) + 1) if marker is not None else None
                    images.append({
                        "sheet": ws.title,
                        "index": idx,
                        "row": row,
                        "column": col,
                        "format": fmt,
                        "data_base64": base64.b64encode(raw).decode("ascii"),
                    })
                except Exception:
                    continue
        wb_img.close()
    except Exception:
        pass
    return images



def _parse_permisos_seguros_sheet(ws):
    """Interpreta la estructura horizontal real de PERMISOS & SEGUROS."""
    import datetime as _dt
    blocks = [
        {"name": "Permiso de navegación en Miami", "start_col": 1, "end_col": 7},
        {"name": "Permiso de navegación en Costa Rica", "start_col": 8, "end_col": 15},
        {"name": "Seguro del INS", "start_col": 16, "end_col": 21},
        {"name": "Extensión de garantía Seakeeper", "start_col": 22, "end_col": 32},
        {"name": "Equipo de salvamento del barco", "start_col": 33, "end_col": 37},
    ]
    records = []
    for b in blocks:
        rec = {"documento": b["name"], "fecha_vencimiento": None,
               "fecha_vencimiento_iso": None, "celda_vencimiento": None,
               "start_col": b["start_col"], "end_col": b["end_col"],
               "observaciones": []}
        dates = []
        for r in range(1, 5):
            for c in range(b["start_col"], b["end_col"] + 1):
                v = ws.cell(r, c).value
                if isinstance(v, (_dt.datetime, _dt.date)):
                    dates.append((r, c, v.date() if isinstance(v, _dt.datetime) else v))
                elif isinstance(v, str):
                    m = re.search(r"\b(\d{1,2})[/-](\d{1,2})[/-](20\d{2})\b", v)
                    if m:
                        try:
                            dates.append((r, c, _dt.date(int(m.group(3)), int(m.group(1)), int(m.group(2)))))
                        except Exception:
                            pass
        labeled = []
        for r in range(1, 5):
            for c in range(b["start_col"], b["end_col"] + 1):
                txt = norm(str(ws.cell(r, c).value or ""))
                if "vence" in txt or "vencimiento" in txt or "expir" in txt:
                    near = [x for x in dates if abs(x[0]-r) <= 2 and abs(x[1]-c) <= 2]
                    labeled.extend(near)
        chosen = labeled[0] if labeled else (dates[0] if dates else None)
        if chosen:
            rr, cc, dt = chosen
            rec["fecha_vencimiento"] = dt.isoformat()
            rec["fecha_vencimiento_iso"] = dt.isoformat()
            rec["celda_vencimiento"] = ws.cell(rr, cc).coordinate
        for r in range(1, 4):
            for c in range(b["start_col"], b["end_col"] + 1):
                txt = str(ws.cell(r, c).value or "").strip()
                if txt and norm(txt) not in {"vence", "anual"}:
                    rec["observaciones"].append(f"{ws.cell(r,c).coordinate}: {txt}")
        records.append(rec)
    return records


def _associate_permit_images(data):
    """Relaciona cada foto con el bloque documental por su columna de anclaje."""
    blocks = [
        ("Permiso de navegación en Miami", 1, 7),
        ("Permiso de navegación en Costa Rica", 8, 15),
        ("Seguro del INS", 16, 21),
        ("Extensión de garantía Seakeeper", 22, 32),
        ("Equipo de salvamento del barco", 33, 37),
    ]
    for im in data.get("images", []):
        if norm(im.get("sheet", "")) != "permisos & seguros":
            continue
        col = int(im.get("column") or 0)
        for name, start, end in blocks:
            if start <= col <= end:
                im["documento_asociado"] = name
                break
    return data


def build_from_excel(path: Path):
    """Construye la base directamente desde Excel sin guardar una copia descomprimida.

    openpyxl abre el .xlsx como archivo ZIP y read_only evita cargar el libro
    completo en memoria de una sola vez. Esto permite trabajar mejor con
    archivos grandes y variables.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    data = {
        "metadata": {
            "agent": "Agente Tiara",
            "version": "2.2",
            "source_of_truth": "Excel cargado por el usuario",
            "source_file": path.name,
        },
        "sheets": {},
        "maintenance_records": [],
        "log_records": [],
        "inventory_records": [],
        "fuel_records": [],
        "budget_records": [],
        "invoice_records": [],
        "cleaning_records": [],
        "documents_records": [],
        "permit_records": [],
    }

    # Conservamos la estructura de hojas para que el resto de los agentes
    # siga funcionando, pero usamos lectura secuencial.
    for ws in wb.worksheets:
        rows = _read_sheet_rows(ws)
        data["sheets"][ws.title] = {
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "rows": rows,
        }

    # Mantenimiento
    ws = _find_sheet(wb, "CHECKLIST MANT. PREVENTIVO")
    if ws:
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [serialize(v) for v in row]
            if any(v not in (None, "") for v in vals):
                data["maintenance_records"].append({
                    "source_row": r,
                    "data": {
                        "mantenimiento": vals[0] if len(vals) > 0 else None,
                        "cada_horas": vals[1] if len(vals) > 1 else None,
                        "historial_horas": vals[2] if len(vals) > 2 else None,
                        "fecha": vals[8] if len(vals) > 8 else None,
                        "observaciones": vals[9] if len(vals) > 9 else None,
                    },
                })

    mapping = {
        "log_records": "VITACORA BARCO EN LINEA",
        "inventory_records": "INVENTARIO",
        "fuel_records": "CONSUMO COMBUSTIBLE ",
        "budget_records": "PRESUPUESTO TIARA 2026",
        "invoice_records": "FACTURAS 2026...",
        "cleaning_records": "S.LIMPIEZA & OTROS",
        "documents_records": "PERMISOS & SEGUROS",
    }
    for key, sheet_name in mapping.items():
        ws = _find_sheet(wb, sheet_name)
        if not ws:
            continue
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [serialize(v) for v in row]
            if not any(v not in (None, "") for v in vals):
                continue

            if key == "log_records":
                # La bitácora tiene una estructura propia que el agente debe
                # entender, no tratar como una fila de texto genérica.
                data[key].append(parse_log_record(vals, r))
            else:
                data[key].append({"source_row": r, "values": vals})

    ws_permits = _find_sheet(wb, "PERMISOS & SEGUROS")
    if ws_permits:
        data["permit_records"] = _parse_permisos_seguros_sheet(ws_permits)

    wb.close()
    # Segunda pasada únicamente para fotografías incrustadas; read_only no las expone.
    # Se analizan con visión y luego se eliminan los bytes para que Supabase no
    # reciba el contenido binario de todas las fotos.
    raw_images = _extract_excel_images(path)
    data["images"], vision_meta = _analyze_embedded_images(raw_images)
    data = _associate_permit_images(data)
    data["metadata"]["vision"] = vision_meta
    return data


def build_from_uploaded_excel(uploaded_file):
    """Procesa un Excel subido desde la aplicación sin descomprimirlo manualmente.

    El formato .xlsx/.xlsm ya es un contenedor ZIP. openpyxl lo abre
    directamente y read_only permite leerlo por partes. Primero copiamos el
    archivo a un temporal por bloques y comprobamos que el contenedor sea
    válido; después construimos la misma base que usa el agente.
    """
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix not in (".xlsx", ".xlsm"):
        raise ValueError("El archivo debe ser Excel .xlsx o .xlsm")

    # Límite de seguridad razonable para evitar cargar accidentalmente un
    # archivo enorme. El límite de Streamlit también puede existir en el
    # servidor, pero este control ocurre dentro de la aplicación.
    max_bytes = 500 * 1024 * 1024
    declared_size = getattr(uploaded_file, "size", None)
    if declared_size and declared_size > max_bytes:
        raise ValueError("El Excel supera el límite de 500 MB de esta aplicación.")

    h = hashlib.sha256()

    with tempfile.NamedTemporaryFile(
        prefix="tiara_excel_",
        suffix=suffix,
        delete=False,
    ) as tmp:
        temp_path = Path(tmp.name)

        while True:
            chunk = uploaded_file.read(1024 * 1024)
            if not chunk:
                break
            h.update(chunk)
            tmp.write(chunk)

    uploaded_file.seek(0)

    try:
        # Un XLSX/XLSM válido es un ZIP. No extraemos el contenido a una
        # carpeta: openpyxl lo lee directamente desde el contenedor.
        if not zipfile.is_zipfile(temp_path):
            raise ValueError(
                "El archivo no parece ser un Excel .xlsx/.xlsm válido. "
                "Vuelve a cargar el archivo original."
            )

        data = build_from_excel(temp_path)

        # Metadatos de la carga actual.
        data["metadata"]["upload_sha256"] = h.hexdigest()
        data["metadata"]["version"] = APP_VERSION
        data["metadata"]["source_file"] = uploaded_file.name
        data["metadata"]["source_size_mb"] = round(
            temp_path.stat().st_size / (1024 * 1024), 2
        )
        data["metadata"]["source_updated_in_app"] = datetime.now(APP_TIMEZONE).isoformat(timespec="seconds")

        # Verificación mínima: el libro debe tener hojas.
        if not data.get("sheets"):
            raise ValueError("El Excel no contiene hojas de datos legibles.")

        return data

    finally:
        try:
            temp_path.unlink(missing_ok=True)
        except Exception:
            pass

def serialize(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


def sheet_rows(data, name):
    return data.get("sheets", {}).get(name, {}).get("rows", [])


def maintenance_records(data):
    return [r for r in data.get("maintenance_records", []) if r.get("data", {}).get("mantenimiento")]


# ============================================================
# CEREBRO / PLANIFICADOR
# ============================================================

def detect_system(p):
    p = norm(p)
    if any(x in p for x in ("generador", "cummins", "onan")): return "Generador Cummins Onan"
    if any(x in p for x in ("pata", "patas", "outdrive", "muelas")): return "Patas Mercury"
    if any(x in p for x in ("motor", "motores", "mercury", "estribor", "babor")): return "Motores Mercury"
    if any(x in p for x in ("combustible", "gasolina", "diesel", "diésel")): return "Combustible"
    if any(x in p for x in ("inventario", "existencias", "stock", "producto", "vasos", "fender", "chaleco")): return "Inventario"
    if any(x in p for x in ("factura", "facturas", "invoice")): return "Facturas"
    if any(x in p for x in ("permiso", "permisos", "seguro", "seguros", "garantia", "garantía", "documento")): return "Permisos y seguros"
    if any(x in p for x in ("bitacora", "bitácora", "evento", "problema", "falla", "alarma", "reparacion", "reparación")): return "Bitácora"
    if any(x in p for x in ("limpieza", "limpiar", "aseo")): return "Limpieza"
    if any(x in p for x in ("presupuesto", "gasto", "gastamos", "costo", "costos", "gastado", "gastó")): return "Finanzas"
    return "General"


def detect_intent(p):
    p = norm(p)
    if any(x in p for x in ("estado general", "estado del barco", "estado completo", "revision general", "revisión general")): return "estado_general"

    # Finanzas: primero identificar un rubro, pero separar INVENTARIO operativo
    # de INVENTARIO.COCINA del presupuesto.
    finance_words = ("cuanto", "gasto", "gaste", "gastamos", "costo", "presupuesto", "presupuestado", "gastado", "gastó", "cuanto se gasto")
    rubric = identify_rubric(p)
    operational_inventory = any(x in p for x in ("que hay en inventario", "qué hay en inventario", "que tenemos en inventario", "existencias", "stock"))
    if rubric and any(x in p for x in finance_words) and not operational_inventory:
        return "gasto_mensual" if (find_month(p) or any(x in p for x in ("ese mes", "mismo mes"))) else "gasto_anual"
    if any(x in p for x in ("cuanto gastamos", "cuánto gastamos", "cuanto gasto", "cuánto gasto", "gasto en", "gastamos en", "presupuesto", "gastado")):
        return "gasto_mensual" if find_month(p) else "gasto_anual"

    if any(x in p for x in ("cuantas veces", "cuántas veces", "cuantos cambios", "cuántos cambios", "cantidad de cambios", "cantidad de mantenimientos")): return "cantidad_mantenimiento"
    if any(x in p for x in ("historial completo", "todo el historial", "todas las veces", "todos los cambios", "todos los mantenimientos", "muestrame el historial", "muéstrame el historial")): return "historial_completo"
    if any(x in p for x in ("ultimo cambio", "último cambio", "ultimo mantenimiento", "último mantenimiento", "cuando fue", "cuándo fue", "cuando se hizo", "cuándo se hizo", "ultima vez", "última vez")): return "ultimo_mantenimiento"
    if any(x in p for x in ("kit 300", "kit de 300", "300 horas", "servicio 300")): return "kit_300"
    if any(x in p for x in ("cuando toca", "cuándo toca", "proximo", "próximo", "faltan horas", "que mantenimiento", "qué mantenimiento", "mantenimiento")): return "mantenimiento"
    if operational_inventory or any(x in p for x in ("inventario", "existencias", "stock")): return "inventario"
    if any(x in p for x in ("combustible", "gasolina", "diesel", "diésel")): return "combustible"
    if any(x in p for x in ("factura", "facturas", "invoice")): return "facturas"
    if any(x in p for x in ("foto", "fotografia", "fotografía", "imagen", "documento fotografiado", "que dice la foto", "qué dice la foto", "lee la foto", "leer la foto", "mira la foto", "que muestra la foto", "qué muestra la foto")): return "fotografias"
    if any(x in p for x in ("permiso", "permisos", "seguro", "seguros", "vence", "vencimiento", "documentos")): return "documentos"
    if any(x in p for x in ("limpieza", "limpiar", "aseo")): return "limpieza"
    if any(x in p for x in ("bitacora", "bitácora", "evento", "falla", "alarma", "problema")): return "bitacora"
    if any(x in p for x in ("horas", "horometro", "horómetro")): return "horometros"
    return "general"


def find_month(p):
    for name, number in MONTHS.items():
        if name in p:
            return number
    return None


def find_years(p):
    """Extrae todos los años explícitos, conservando el orden de aparición."""
    years = []
    for raw in re.findall(r"\b(20\d{2})\b", p):
        year = int(raw)
        if year not in years:
            years.append(year)
    return years


def find_year(p):
    years = find_years(p)
    return years[0] if years else None


def detect_financial_operation(p):
    """Detecta únicamente operaciones financieras aprobadas por el proyecto."""
    p = norm(p)
    # Operación específica de SALIDAS BARCO: únicamente combustible dividido
    # entre el número de salidas, expresado en USD por salida.
    if any(x in p for x in (
        "combustible por salida", "combustible por cada salida",
        "gasto de combustible por salida", "gasto combustible por salida",
        "costo de combustible por salida", "costo combustible por salida",
        "promedio de combustible por salida", "promedio combustible por salida",
        "combustible promedio por salida",
        "usd por salida", "dolares por salida", "dolares por cada salida",
        "combustible / salidas", "combustible entre salidas",
        "combustible dividido entre salidas", "costo promedio por salida",
    )):
        return "costo_combustible_por_salida"

    # El orden importa: una pregunta de variación/comparación no debe quedar
    # clasificada como una simple consulta de presupuesto.
    if any(x in p for x in (
        "comparar", "compara", "comparacion", "comparación", "versus", " vs ",
        "diferencia", "diferencias", "variacion", "variación", "porcentaje",
        "%", "por ciento", "aumento", "aumentó", "incremento", "incrementó",
        "disminucion", "disminución", "disminuyo", "disminuyó", "subio", "subió",
        "bajo", "bajó", "crecio", "creció", "cuanto mas", "cuánto más",
        "cuanto menos", "cuánto menos", "mismo periodo", "mismo mes", "misma fecha"
    )):
        return "comparacion"
    if any(x in p for x in ("promedio", "promedios", "media")):
        return "promedio"
    if any(x in p for x in (
        "gasto vs presupuesto", "gasto versus presupuesto",
        "gastado vs presupuestado", "gastado versus presupuestado",
        "contra el presupuesto", "contra presupuesto", "vs presupuesto",
        "versus presupuesto", "excedio el presupuesto", "excedió el presupuesto",
        "por encima del presupuesto", "por debajo del presupuesto",
        "presupuesto contra gasto", "gasto contra presupuesto"
    )):
        return "gasto_vs_presupuesto"
    if any(x in p for x in (
        "consumido vs gastado", "consumido versus gastado",
        "consumo vs gasto", "consumo versus gasto",
        "consumido y gastado", "consumo y gasto",
        "litros vs dolares", "litros versus dolares",
        "litros vs gasto", "litros versus gasto"
    )):
        return "consumido_vs_gastado"
    return None


def think(question):
    p = norm(question)
    years = find_years(p)
    return {
        "intent": detect_intent(p),
        "system": detect_system(p),
        "month": find_month(p),
        "year": years[0] if years else None,
        "years": years,
        "rubric": identify_rubric(p),
        "financial_operation": detect_financial_operation(p),
        "normalized": p,
        "original_question": question,
    }


# ============================================================
# MÓDULO DE MANTENIMIENTO
# ============================================================

SYSTEM_ALIASES = {
    "Motores Mercury": ["motor", "motores", "mercury"],
    "Patas Mercury": ["pata", "patas"],
    "Generador Cummins Onan": ["generador", "cummins", "onan"],
}


def system_matches(name, system):
    n = norm(name)
    return any(a in n for a in SYSTEM_ALIASES.get(system, []))


def maintenance_candidates(data, system, p):
    records = maintenance_records(data)
    exact = []
    broad = []
    for r in records:
        d = r["data"]
        name = norm(d.get("mantenimiento"))
        row = int(r.get("source_row", 0) or 0)

        if system == "Motores Mercury":
            # Bloque de motores: filas 4-9 del checklist.
            if 4 <= row <= 9:
                broad.append(r)
                if "aceite motores" in name or "filtros de aceite" in name:
                    exact.append(r)

        elif system == "Patas Mercury":
            # Las patas aparecen dentro del bloque de motores, pero se
            # identifican por su nombre específico.
            if "pata" in name:
                broad.append(r)
                if "aceite de pata" in name:
                    exact.append(r)

        elif system == "Generador Cummins Onan":
            # Bloque del generador: filas 11-31 del checklist.
            if 11 <= row <= 31:
                broad.append(r)
                if "cambio de aceite" in name or "cambio filtro de aceite" in name:
                    exact.append(r)

    return exact, broad


def maintenance_row_summary(r):
    d = r["data"]
    hours = extract_hours(d.get("historial_horas"))
    dates = extract_dates(d.get("fecha"))
    interval = safe_float(d.get("cada_horas"))
    return {
        "name": str(d.get("mantenimiento") or "").strip(),
        "interval": int(interval) if interval is not None and interval.is_integer() else interval,
        "hours": hours,
        "dates": dates,
        "last_hour": hours[-1] if hours else None,
        "last_date": dates[-1] if dates else None,
        "next_hour": (hours[-1] + int(interval)) if hours and interval is not None else None,
        "observations": str(d.get("observaciones") or "").strip(),
    }


def answer_ultimo_mantenimiento(data, thought):
    system = thought["system"]
    p = thought["normalized"]
    if system not in SYSTEM_ALIASES:
        return "No pude determinar si preguntas por motores, patas o generador."

    exact, broad = maintenance_candidates(data, system, p)
    wants_oil = "aceite" in p or "filtro" in p
    candidates = exact if wants_oil and exact else broad

    # Si es una pregunta de motores sin detalle, devolver SOLO el mantenimiento principal,
    # no todas las filas relacionadas.
    if not wants_oil and system == "Motores Mercury":
        candidates = [r for r in broad if "aceite motores" in norm(r["data"].get("mantenimiento"))]
    elif not wants_oil and system == "Patas Mercury":
        # Para una pregunta genérica por las patas, usar el mantenimiento
        # principal (aceite), no escoger arbitrariamente un ánodo u otro ítem.
        candidates = [r for r in broad if "aceite de pata" in norm(r["data"].get("mantenimiento"))]

    if not candidates:
        return "No encontré un registro de mantenimiento confirmado para ese sistema."

    # Escoger la fila más reciente por fecha; en empate, por horómetro.
    summaries = [maintenance_row_summary(r) for r in candidates]
    summaries.sort(key=lambda x: (x["last_date"] or date.min, x["last_hour"] or -1))
    s = summaries[-1]

    lines = [
        f"### Último mantenimiento — {system}",
        f"**Mantenimiento:** {s['name']}",
        f"**Última fecha registrada:** {fmt_date(s['last_date'])}",
        f"**Último horómetro registrado:** {s['last_hour']} h" if s["last_hour"] is not None else "**Último horómetro registrado:** NO DETERMINADO",
    ]
    if s["interval"] is not None and s["last_hour"] is not None:
        lines.append(f"**Próximo por horas:** {s['next_hour']} h")
    return "\n".join(lines)


def answer_historial(data, thought, complete=False):
    system = thought["system"]
    if system not in SYSTEM_ALIASES:
        return "No pude determinar el sistema para consultar el historial."
    exact, broad = maintenance_candidates(data, system, thought["normalized"])
    rows = broad if complete else exact
    if not rows:
        rows = broad
    if not rows:
        return "No encontré historial de mantenimiento para ese sistema."

    # Historial completo de un sistema: máximo los registros que realmente pertenecen al sistema.
    items = []
    for r in rows:
        s = maintenance_row_summary(r)
        if s["last_date"] or s["last_hour"]:
            items.append(s)
    items.sort(key=lambda x: (x["last_date"] or date.min, x["last_hour"] or -1), reverse=True)

    lines = [f"### Historial de mantenimiento — {system}"]
    for s in items:
        if s["last_date"]:
            lines.append(f"- **{s['name']}** — última fecha: {fmt_date(s['last_date'])}" + (f", {s['last_hour']} h" if s["last_hour"] is not None else ""))
        else:
            lines.append(f"- **{s['name']}** — última hora: {s['last_hour']} h")
    return "\n".join(lines)


def answer_count(data, thought):
    system = thought["system"]
    exact, broad = maintenance_candidates(data, system, thought["normalized"])
    rows = exact if exact else broad
    if not rows:
        return "No encontré registros suficientes para contar ese mantenimiento."
    total = 0
    details = []
    for r in rows:
        s = maintenance_row_summary(r)
        n = len(s["hours"]) if s["hours"] else len(s["dates"])
        if n:
            total += n
            details.append(f"- {s['name']}: {n} registros")
    if not details:
        return "No encontré registros fechados o por horas para ese mantenimiento."
    return f"### Cantidad de mantenimientos — {system}\n**Total:** {total}\n" + "\n".join(details)


def answer_mantenimiento(data, thought):
    system = thought["system"]
    p = thought["normalized"]
    if system not in SYSTEM_ALIASES:
        return "No pude determinar el sistema del barco."

    if system == "Motores Mercury":
        rows = [r for r in maintenance_records(data) if "aceite motores" in norm(r["data"].get("mantenimiento"))]
    elif system == "Patas Mercury":
        rows = [r for r in maintenance_records(data) if "aceite de pata" in norm(r["data"].get("mantenimiento"))]
    else:
        rows = [r for r in maintenance_records(data) if "cambio de aceite" in norm(r["data"].get("mantenimiento")) and 11 <= int(r.get("source_row", 0)) <= 31]

    if not rows:
        return "No encontré el mantenimiento programado para ese sistema."

    # Si existen varias filas del mismo sistema, tomar la que tenga el
    # registro más reciente, no simplemente la primera fila del Excel.
    summaries = [maintenance_row_summary(r) for r in rows]
    s = max(summaries, key=lambda x: (x["last_date"] or date.min, x["last_hour"] or -1))
    current = current_hours(data).get(system)
    lines = [
        f"### 🔧 Próximo mantenimiento — {system}",
        f"**Mantenimiento:** {s['name']}",
        f"**Último registrado:** {fmt_date(s['last_date'])}",
        f"**Horómetro del último registro:** {s['last_hour']} h" if s["last_hour"] is not None else "**Horómetro del último registro:** NO DETERMINADO",
        f"**Intervalo:** {s['interval']} h" if s["interval"] is not None else "**Intervalo:** NO DETERMINADO",
        f"**Próximo:** {s['next_hour']} h" if s["next_hour"] is not None else "**Próximo:** NO DETERMINADO",
    ]
    if current is not None and s["next_hour"] is not None:
        faltan = s["next_hour"] - current
        lines.append(f"**Horómetro actual:** {current} h")
        lines.append(f"**Faltan:** {faltan} h" if faltan >= 0 else f"**Excedido por:** {abs(faltan)} h")
    return "\n".join(lines)


def answer_kit(data):
    rows = [r for r in maintenance_records(data) if "kit 300 horas" in norm(r["data"].get("mantenimiento"))]
    if not rows:
        return "No encontré el Kit de 300 horas en el Excel."

    summaries = [maintenance_row_summary(r) for r in rows]
    s = max(summaries, key=lambda x: (x["last_date"] or date.min, x["last_hour"] or -1))
    current = current_hours(data).get("Motores Mercury")
    lines = [
        "### 🔩 Kit 300 horas — Motores Mercury",
        f"**Último kit registrado:** {s['last_hour']} h" if s["last_hour"] is not None else "**Último kit registrado:** NO DETERMINADO",
        f"**Fecha del último kit:** {fmt_date(s['last_date'])}",
        f"**Intervalo:** {s['interval']} h" if s["interval"] is not None else "**Intervalo:** NO DETERMINADO",
        f"**Próximo kit:** {s['next_hour']} h" if s["next_hour"] is not None else "**Próximo kit:** NO DETERMINADO",
    ]
    if current is not None and s["next_hour"] is not None:
        faltan = s["next_hour"] - current
        lines.append(f"**Horómetro actual de motores:** {current} h")
        lines.append(f"**Faltan:** {faltan} h" if faltan >= 0 else f"**Excedido por:** {abs(faltan)} h")
    return "\n".join(lines)


def current_hours(data):
    result = {}
    # La bitácora es la fuente más reciente cuando registra un evento con
    # horas. El checklist queda como respaldo histórico.
    fallback_rows = maintenance_records(data)
    targets = {
        "Motores Mercury": 4,
        "Generador Cummins Onan": 11,
    }
    for system, source_row in targets.items():
        log_h = latest_log_hours(data, system)
        if log_h is not None:
            result[system] = int(log_h) if float(log_h).is_integer() else log_h
            continue

        match = next((r for r in fallback_rows if int(r.get("source_row", -1)) == source_row), None)
        if match:
            hours = extract_hours(match["data"].get("historial_horas"))
            result[system] = hours[-1] if hours else None
        else:
            result[system] = None

    # Patas utilizan deliberadamente el mismo horómetro de motores.
    result["Patas Mercury"] = result.get("Motores Mercury")
    return result


def answer_state(data):
    hours = current_hours(data)
    lines = ["### Estado general del barco"]

    target_names = {
        "Motores Mercury": {"cambio aceite motores"},
        "Patas Mercury": {"cambio de aceite de pata motores"},
        "Generador Cummins Onan": {"cambio de aceite"},
    }

    for system in ("Motores Mercury", "Patas Mercury", "Generador Cummins Onan"):
        h = hours.get(system)
        lines.append(
            f"**{system}:** {h} h" if h is not None
            else f"**{system}:** NO DETERMINADO"
        )

        rows = [
            r for r in maintenance_records(data)
            if norm(r["data"].get("mantenimiento")) in target_names[system]
        ]
        if not rows:
            continue

        summaries = [maintenance_row_summary(r) for r in rows]
        s = max(summaries, key=lambda x: (x["last_date"] or date.min, x["last_hour"] or -1))

        if s["last_date"]:
            lines.append(f"  - Último mantenimiento: {fmt_date(s['last_date'])}")
        if s["last_hour"] is not None:
            lines.append(f"  - Horómetro del último mantenimiento: {s['last_hour']} h")
        if s["interval"] is not None:
            lines.append(f"  - Intervalo: {s['interval']} h")
        if s["next_hour"] is not None:
            lines.append(f"  - Próximo: {s['next_hour']} h")
            if h is not None:
                faltan = s["next_hour"] - h
                lines.append(
                    f"  - Faltan: {faltan} h" if faltan >= 0
                    else f"  - Excedido por: {abs(faltan)} h"
                )

    return "\n".join(lines)


# ============================================================
# FINANZAS / COMBUSTIBLE
# ============================================================

# Alias de todos los rubros financieros presentes en la base.
# La consulta puede usar el nombre exacto del Excel o una forma coloquial.
# La lógica devuelve siempre el nombre canónico que usa el presupuesto.
RUBRIC_RULES = [
    (("mantenimiento", "anual"), "MANTENIMIENTO ANUAL TIARA"),
    (("mantenimiento", "motor"), "MANT.MOTORES"),
    (("mantenimiento", "motores"), "MANT.MOTORES"),
    (("motores",), "MANT.MOTORES"),
    (("mantenimiento", "generador"), "MANT.GENERADOR"),
    (("generador",), "MANT.GENERADOR"),
    (("combustible",), "COMSUMO COMBUSTIBLE"),
    (("gasolina",), "COMSUMO COMBUSTIBLE"),
    (("diesel",), "COMSUMO COMBUSTIBLE"),
    (("diésel",), "COMSUMO COMBUSTIBLE"),
    (("repuestos",), "REPUESTOS & COTIZACIONES"),
    (("repuesto",), "REPUESTOS & COTIZACIONES"),
    (("cotizaciones",), "REPUESTOS & COTIZACIONES"),
    (("cotizacion",), "REPUESTOS & COTIZACIONES"),
    (("seguros",), "SEGUROS & MEMBRESIAS"),
    (("seguro",), "SEGUROS & MEMBRESIAS"),
    (("membresias",), "SEGUROS & MEMBRESIAS"),
    (("membresía",), "SEGUROS & MEMBRESIAS"),
    (("salario",), "SALARIO ANDRES"),
    (("sueldo",), "SALARIO ANDRES"),
    (("limpieza",), "S.LIMPIEZA"),
    (("inventario",), "INVENTARIO.COCINA"),
    (("cocina",), "INVENTARIO.COCINA"),
    (("s.cocina",), "INVENTARIO.COCINA"),
    (("agua",), "PAGO POR AGUA DULCE HIELO ETC."),
    (("hielo",), "PAGO POR AGUA DULCE HIELO ETC."),
    (("agua dulce",), "PAGO POR AGUA DULCE HIELO ETC."),
    (("visitas",), "VISITAS JUAN MANUEL"),
    (("juan manuel",), "VISITAS JUAN MANUEL"),
    (("marina",), "TRAVEL/SLIP MARINA"),
    (("slip",), "TRAVEL/SLIP MARINA"),
    (("travel",), "TRAVEL/SLIP MARINA"),
    (("comida",), "COMIDA ANDRES"),
    (("ccss",), "CCSS ANDRES"),
    (("otros",), "OTROS/"),
]


def identify_rubric(p):
    """Reconoce rubros financieros por lenguaje natural y por aliases del Excel."""
    p = norm(p)
    if any(x in p for x in ("que hay en inventario", "qué hay en inventario", "existencias", "stock")) and not any(x in p for x in ("gasto", "gastamos", "costo", "presupuesto", "gastado")):
        return None
    aliases = [
        (("mantenimiento anual", "mantenimiento general anual", "mant anual"), "MANTENIMIENTO ANUAL TIARA"),
        (("mantenimiento de motores", "mantenimiento motores", "mant motores", "motor mercury"), "MANT.MOTORES"),
        (("mantenimiento del generador", "mantenimiento generador", "mant generador"), "MANT.GENERADOR"),
        (("combustible", "gasolina", "diesel", "diésel", "consumo de combustible"), "COMSUMO COMBUSTIBLE"),
        (("repuestos", "repuesto", "cotizaciones", "cotizacion"), "REPUESTOS & COTIZACIONES"),
        (("seguros", "seguro", "membresias", "membresía", "membresias y seguros"), "SEGUROS & MEMBRESIAS"),
        (("salario", "sueldo"), "SALARIO ANDRES"),
        (("limpieza", "aseo"), "S.LIMPIEZA"),
        (("inventario cocina", "inventario de cocina", "s cocina", "s.cocina"), "INVENTARIO.COCINA"),
        (("agua dulce", "hielo", "agua e hielo", "agua"), "PAGO POR AGUA DULCE HIELO ETC."),
        (("visitas", "juan manuel"), "VISITAS JUAN MANUEL"),
        (("marina", "slip", "travel slip", "travel"), "TRAVEL/SLIP MARINA"),
        (("comida", "comidas", "alimentacion", "alimentación"), "COMIDA ANDRES"),
        (("ccss",), "CCSS ANDRES"),
        (("otros", "otros gastos"), "OTROS/"),
        (("salidas barco", "salida barco", "salidas de barco", "salida de barco", "salidas del barco", "salida del barco"), "SALIDAS BARCO"),
    ]
    for variants, rubric in aliases:
        if any(v in p for v in variants):
            return rubric
    ordered = sorted(RUBRIC_RULES, key=lambda item: len(item[0]), reverse=True)
    for words, rubric in ordered:
        if all(norm(w) in p for w in words): return rubric
    return None


def _budget_name(rec):
    data = rec.get("data") or {}
    if data:
        for key in ("RUBROS", "rubros", "RUBRO", "rubro"):
            if data.get(key) not in (None, ""):
                return str(data[key]).strip()
    vals = rec.get("values") or []
    if len(vals) > 2 and vals[2] not in (None, ""):
        return str(vals[2]).strip()
    return ""


def _budget_number(value):
    if value in (None, "", "-", "?"):
        return None
    return safe_float(value)


def _budget_month_key(year, month):
    return f"{MONTH_NAMES[month][:3]}-{str(year)[2:]}"


def _budget_dict_value(rec, year, month):
    """Lee un mes desde la estructura JSON real del presupuesto."""
    data = rec.get("data") or {}
    if not data:
        return None

    target = _budget_month_key(year, month)
    for key, value in data.items():
        nk = norm(key).replace(" ", "")
        # Acepta ene-26, ene26, feb-26 y también ' feb 26'.
        if nk.replace("-", "") == target.replace("-", ""):
            return _budget_number(value)
    return None


def _budget_rows_from_excel_values(data):
    """Convierte filas antiguas del Excel en registros con encabezados reales."""
    rows = budget_rows(data)
    headers = None
    raw = sheet_rows(data, "PRESUPUESTO TIARA 2026")
    if raw:
        # En el Excel del proyecto, la fila 3 (índice 2) contiene los encabezados.
        if len(raw) > 2:
            headers = [norm(x).replace(" ", "") if x is not None else "" for x in raw[2]]

    out = []
    for rec in rows:
        if rec.get("data"):
            out.append(rec)
            continue
        vals = rec.get("values") or []
        mapped = {}
        if headers:
            for i, value in enumerate(vals):
                if i < len(headers) and headers[i]:
                    mapped[headers[i]] = value
        clone = dict(rec)
        clone["data"] = mapped
        out.append(clone)
    return out


def budget_rows(data):
    rows = data.get("budget_records", [])
    if rows:
        return rows
    raw = sheet_rows(data, "PRESUPUESTO TIARA 2026")
    return [{"source_row": i + 1, "values": row} for i, row in enumerate(raw)]


def _is_budget_total_row(name, rec=None):
    n = norm(name)
    if n in ("total", "totales") or n.startswith("total "):
        return True
    # En la estructura real del Excel, la fila TOTAL usa "Total" en
    # OBSERVACIONES y deja RUBROS vacío.
    d = (rec or {}).get("data") or {}
    for key in ("OSERVACIONES", "OBSERVACIONES", "observaciones"):
        if norm(d.get(key)) in ("total", "totales"):
            return True
    return False


def _budget_name_matches(name, rubric):
    """Compara nombres del Excel tolerando tildes, espacios y variantes."""
    a = norm(name)
    b = norm(rubric)
    variants = {
        "comsumo combustible": {"comsumo combustible", "consumo combustible"},
        "repuestos & cotizaciones": {"repuestos & cotizaciones", "repuestos &cotizaciones"},
        "s.limpieza": {"s.limpieza", "s limpieza"},
        "seguros & membresias": {"seguros & membresias", "seguros/ membresias", "seguros/membresias"},
        "pago por agua dulce hielo etc.": {
            "pago por agua dulce hielo etc.",
            "pago por agua dulce,hielo,..",
        },
    }
    if a == b:
        return True
    return a in variants.get(b, set()) or b in variants.get(a, set())


def _find_budget_record(data, rubric):
    target = norm(rubric)
    for rec in _budget_rows_from_excel_values(data):
        if _budget_name_matches(_budget_name(rec), rubric):
            return rec
    return None


def _find_budget_record_for_month(data, rubric, year, month):
    """Busca el registro del rubro que realmente contiene el mes solicitado.

    El Excel puede contener más de una fila con el mismo rubro (por ejemplo,
    una fila de referencia/presupuesto y otra con los importes mensuales).
    Para una pregunta mensual no debemos quedarnos con la primera fila si esa
    fila no tiene el mes solicitado.
    """
    target = norm(rubric)
    candidates = []
    for rec in _budget_rows_from_excel_values(data):
        if not _budget_name_matches(_budget_name(rec), rubric):
            continue
        value = _budget_dict_value(rec, year, month)
        if value is not None:
            candidates.append((rec, value))
    return candidates[0] if candidates else (None, None)


def budget_value_for_month(data, rubric, year, month):
    rec = _find_budget_record(data, rubric)
    if not rec:
        return None

    value = _budget_dict_value(rec, year, month)
    if value is not None:
        return value

    # Compatibilidad con el formato antiguo de filas del Excel.
    vals = rec.get("values", [])
    if year == 2026:
        col = 3 + (month - 1) * 2
    elif year == 2025:
        col = 4 + (month - 1) * 2
    elif year == 2024:
        col = 3 + (month - 1) * 2
    else:
        return None
    return _budget_number(vals[col]) if col < len(vals) else None


def budget_total(data, rubric, year):
    rec = _find_budget_record(data, rubric)
    if not rec:
        return None
    d = rec.get("data") or {}
    for key in (f"TOTAL  {year}", f"TOTAL {year}"):
        if key in d:
            return _budget_number(d[key])
    # Buscar de forma tolerante.
    for key, value in d.items():
        if norm(key).replace(" ", "") == f"total{year}":
            return _budget_number(value)

    vals = rec.get("values", [])
    if year == 2026 and len(vals) > 27:
        return _budget_number(vals[27])
    if year == 2025 and len(vals) > 28:
        return _budget_number(vals[28])
    return None


def budget_total_all(data, year, month=None):
    """Total general del presupuesto. Usa la fila TOTAL del Excel cuando existe."""
    rows = _budget_rows_from_excel_values(data)

    # 1) La fuente tiene una fila explícita TOTAL. Es la autoridad para una
    #    pregunta como 'cuánto se gastó en agosto de 2026'.
    for rec in rows:
        if _is_budget_total_row(_budget_name(rec), rec):
            if month is not None:
                value = _budget_dict_value(rec, year, month)
                if value is not None:
                    return value
            else:
                d = rec.get("data") or {}
                for key, value in d.items():
                    if norm(key).replace(" ", "") == f"total{year}":
                        return _budget_number(value)

    # 2) Compatibilidad: sumar los rubros si la fila TOTAL no está disponible.
    total = 0.0
    found = False
    for rec in rows:
        name = _budget_name(rec)
        if not name or _is_budget_total_row(name, rec):
            continue
        value = _budget_dict_value(rec, year, month) if month else budget_total(data, name, year)
        if value is not None:
            total += value
            found = True
    return total if found else None


def _financial_values_for_period(data, rubric, years, month=None):
    """Obtiene valores del rubro financiero para los años solicitados.

    Si se indica mes, usa ese mes. Si no, usa el total anual.
    Nunca convierte un dato ausente en cero.
    """
    values = {}
    for year in years:
        value = budget_value_for_month(data, rubric, year, month) if month else budget_total(data, rubric, year)
        values[year] = value
    return values


def _financial_change(old_value, new_value):
    if old_value is None or new_value is None:
        return None, None
    difference = new_value - old_value
    if old_value == 0:
        return difference, None
    return difference, (difference / old_value) * 100.0


def _financial_period_label(year, month=None):
    if month:
        return f"{MONTH_NAMES[month].capitalize()} {year}"
    return str(year)


def _financial_reference_years(thought):
    years = list(thought.get("years") or [])
    if len(years) >= 2:
        return years[0], years[1]
    year = thought.get("year") or 2026
    # Para una comparación sin dos años explícitos, comparar el año solicitado
    # contra el año inmediatamente anterior.
    return year - 1, year


FINANCIAL_SOURCE_MAP = {
    "COMSUMO COMBUSTIBLE": "CONSUMO COMBUSTIBLE ",
    "S.LIMPIEZA": "S.LIMPIEZA & OTROS",
    "INVENTARIO.COCINA": "INVENTARIO",
    "MANT.MOTORES": "VITACORA BARCO EN LINEA",
    "MANT.GENERADOR": "VITACORA BARCO EN LINEA",
    "MANTENIMIENTO ANUAL TIARA": "VITACORA BARCO EN LINEA",
    "REPUESTOS & COTIZACIONES": "FACTURAS 2026...",
    "SEGUROS & MEMBRESIAS": "PERMISOS & SEGUROS",
    "PAGO POR AGUA DULCE HIELO ETC.": "FACTURAS 2026...",
    "VISITAS JUAN MANUEL": "S.LIMPIEZA & OTROS",
    "TRAVEL/SLIP MARINA": "FACTURAS 2026...",
    "COMIDA ANDRES": "FACTURAS 2026...",
    "CCSS ANDRES": "FACTURAS 2026...",
    "SALARIO ANDRES": None,
    "OTROS/": "FACTURAS 2026...",
}

def _record_values(rec):
    vals = rec.get("values")
    if isinstance(vals, (list, tuple)):
        return list(vals)
    d = rec.get("data") or {}
    return list(d.values()) if isinstance(d, dict) else []

def _record_data(rec):
    d = rec.get("data")
    return d if isinstance(d, dict) else {}

def _date_from_any(value):
    return parse_date(value) or excel_serial_to_date(value)

def _record_month_year(rec):
    d = _record_data(rec)
    # Primero busca una fecha explícitamente etiquetada.
    for key, value in d.items():
        if "fecha" in norm(key) or norm(key) in {"mes", "month"}:
            dt = _date_from_any(value)
            if dt:
                return dt.year, dt.month
    vals = _record_values(rec)
    # Estructuras conocidas de las hojas financieras.
    sheet = norm(rec.get("source_sheet", ""))
    if "limpieza" in sheet and len(vals) > 10:
        dt = _date_from_any(vals[10])
        if dt:
            return dt.year, dt.month
    if "combustible" in sheet and vals:
        dt = _date_from_any(vals[0])
        if dt:
            return dt.year, dt.month
    # Fallback: buscar cualquier fecha serial/ISO razonable.
    for value in vals:
        dt = _date_from_any(value)
        if dt and 2000 <= dt.year <= 2100:
            return dt.year, dt.month
    return None, None

def _record_money(rec):
    d = _record_data(rec)
    # No usar columnas llamadas VALOR U como gasto: se necesita el total.
    preferred = ("TOTAL", "PRECIO TOAL", "PRECIO TOTAL", "TOTAL FACTURA", "FACTURA", "MONTO", "IMPORTE")
    for key in preferred:
        for k, value in d.items():
            if norm(k) == norm(key):
                n = _budget_number(value)
                if n is not None:
                    return n
    vals = _record_values(rec)
    sheet = norm(rec.get("source_sheet", ""))
    # S.LIMPIEZA & OTROS: PRECIO TOTAL es la octava columna.
    if "limpieza" in sheet and len(vals) > 7:
        return _budget_number(vals[7])
    return None

def _source_records(data, sheet_name):
    target = norm(sheet_name).replace(" ", "")
    rows = []
    for key, value in (data.get("sheets") or {}).items():
        if norm(key).replace(" ", "") != target:
            continue
        if isinstance(value, dict):
            rows = value.get("records") or value.get("rows") or []
        break
    if rows:
        return rows
    # Fallback a los módulos normalizados.
    mapping = {
        "CONSUMO COMBUSTIBLE": data.get("fuel_records", []),
        "S.LIMPIEZA & OTROS": data.get("cleaning_records", []),
        "INVENTARIO": data.get("inventory_records", []),
        "VITACORA BARCO EN LINEA": data.get("log_records", []),
        "FACTURAS 2026...": data.get("invoice_records", []),
        "PERMISOS & SEGUROS": data.get("documents_records", []),
    }
    return mapping.get(sheet_name, [])

def _actual_financial_value(data, rubric, year, month=None):
    """Busca gasto/consumo real en la hoja específica, sin confundir ceros de plantilla con ausencia."""
    sheet = FINANCIAL_SOURCE_MAP.get(rubric)
    if not sheet:
        return None
    records = _source_records(data, sheet)
    if rubric == "COMSUMO COMBUSTIBLE":
        total = 0.0; found = False
        for rec in records:
            d = _record_data(rec)
            dt = _date_from_any(d.get("MES")) or _date_from_any(d.get("FECHA"))
            if not dt:
                vals = _record_values(rec)
                dt = _date_from_any(vals[0]) if vals else None
            if not dt or dt.year != year or (month and dt.month != month):
                continue
            value = None
            for key in ("TOTAL", "FACTURA"):
                value = _budget_number(d.get(key))
                if value is not None:
                    break
            if value is not None:
                total += value; found = True
        return total if found else None

    total = 0.0; found = False
    for rec in records:
        ry, rm = _record_month_year(rec)
        if ry != year or (month is not None and rm != month):
            continue
        value = _record_money(rec)
        if value is not None:
            total += value; found = True
    return total if found else None

def _financial_ytd(data, rubric, year, through_month):
    vals = [_actual_financial_value(data, rubric, year, m) for m in range(1, through_month + 1)]
    valid = [v for v in vals if v is not None]
    return (sum(valid), len(valid)) if valid else (None, 0)

def _salidas_barco_record(data):
    """Busca exclusivamente la fila real de SALIDA(S) BARCO.

    El Excel fuente usa el nombre singular ``SALIDA BARCO``, mientras que
    el cerebro utiliza el nombre canónico ``SALIDAS BARCO``. Esta tolerancia
    se limita a este módulo para no alterar el buscador general de rubros.
    """
    rec = _find_budget_record(data, "SALIDAS BARCO")
    if rec:
        return rec

    for candidate in _budget_rows_from_excel_values(data):
        name = norm(_budget_name(candidate))
        if name in {"salida barco", "salidas barco"}:
            return candidate
    return None


def _salidas_barco_value(data, year, month=None):
    """Obtiene SALIDAS BARCO directamente de su fila real del presupuesto.

    No modifica el lector general de rubros. Para 2026 y 2025 usa los
    encabezados mensuales reales del presupuesto y su TOTAL anual.
    """
    rec = _salidas_barco_record(data)
    if not rec:
        return None

    if month is not None:
        value = _budget_dict_value(rec, year, month)
        if value is not None:
            return value

        # Compatibilidad con la fila antigua ``values`` del Excel.
        vals = rec.get("values", [])
        if year == 2026:
            col = 3 + (month - 1) * 2
        elif year == 2025:
            col = 4 + (month - 1) * 2
        else:
            return None
        return _budget_number(vals[col]) if col < len(vals) else None

    # Primero, TOTAL anual de la propia fila SALIDA BARCO.
    d = rec.get("data") or {}
    for key, value in d.items():
        if norm(key).replace(" ", "") == f"total{year}":
            number = _budget_number(value)
            if number is not None:
                return number

    vals = rec.get("values", [])
    if year == 2026 and len(vals) > 27:
        number = _budget_number(vals[27])
        if number is not None:
            return number
    if year == 2025 and len(vals) > 28:
        number = _budget_number(vals[28])
        if number is not None:
            return number

    # Último respaldo: sumar únicamente los meses numéricos de la propia fila.
    monthly = []
    for m in range(1, 13):
        value = _budget_dict_value(rec, year, m)
        if value is not None:
            monthly.append(value)
    return sum(monthly) if monthly else None


def _period_specs_for_salidas(question, thought):
    """Extrae uno o dos períodos explícitos para SALIDAS BARCO."""
    p = norm(question)
    month_hits = []
    for name, number in MONTHS.items():
        for m in re.finditer(rf"\b{re.escape(name)}\b", p):
            month_hits.append((m.start(), number))
    month_hits.sort()
    year_hits = [(m.start(), int(m.group(1)))
                 for m in re.finditer(r"\b(20\d{2})\b", p)]

    if len(month_hits) >= 2:
        months = [x[1] for x in month_hits[:2]]
        if len(year_hits) >= 2:
            return [(year_hits[0][1], months[0]), (year_hits[1][1], months[1])]
        year = year_hits[0][1] if year_hits else (thought.get("year") or 2026)
        return [(year, months[0]), (year, months[1])]

    if len(month_hits) == 1:
        month = month_hits[0][1]
        years = [y for _, y in year_hits]
        if len(years) >= 2:
            return [(years[0], month), (years[1], month)]
        year = years[0] if years else (thought.get("year") or 2026)
        return [(year, month)]

    years = list(thought.get("years") or [])
    if len(years) >= 2:
        return [(years[0], None), (years[1], None)]
    return [(thought.get("year") or 2026, None)]


def _combustible_budget_for_salidas(data, year, month=None):
    """Obtiene combustible exclusivamente de PRESUPUESTO TIARA 2026."""
    records = _budget_rows_from_excel_values(data)
    matches = [r for r in records if _budget_name_matches(_budget_name(r), "COMSUMO COMBUSTIBLE") or _budget_name_matches(_budget_name(r), "CONSUMO COMBUSTIBLE")]
    if not matches:
        return None
    salidas = _salidas_barco_record(data)
    sr = int(salidas.get("source_row") or 10**9) if salidas else 10**9
    prior = [r for r in matches if int(r.get("source_row") or 0) < sr]
    rec = max(prior, key=lambda r: int(r.get("source_row") or 0)) if prior else matches[-1]
    if month is not None:
        return _budget_dict_value(rec, year, month)
    for key, value in (rec.get("data") or {}).items():
        if norm(key).replace(" ", "") == f"total{year}":
            return _budget_number(value)
    return None


def _combustible_por_salida_result(data, year, month=None):
    fuel = _combustible_budget_for_salidas(data, year, month)
    outings = _salidas_barco_value(data, year, month)
    if fuel is None or outings is None or outings <= 0:
        return None, fuel, outings
    return fuel / outings, fuel, outings


def _answer_combustible_por_salida(data, question, thought):
    """Único cruce financiero aprobado: combustible ÷ salidas."""
    specs = _period_specs_for_salidas(question, thought)
    results = []
    for year, month in specs[:2]:
        cost, fuel, outings = _combustible_por_salida_result(data, year, month)
        results.append((year, month, cost, fuel, outings))

    def label(year, month):
        return _financial_period_label(year, month)

    if len(results) == 1:
        year, month, cost, fuel, outings = results[0]
        return "\n".join([
            "### Análisis financiero — Combustible por salida",
            f"**Período:** {label(year, month)}",
            f"**Gasto de combustible:** {money(fuel) if fuel is not None else 'NO DETERMINADO'}",
            f"**Salidas de barco:** {int(outings) if outings is not None and float(outings).is_integer() else outings if outings is not None else 'NO DETERMINADO'}",
            f"**Combustible promedio por salida:** {money(cost) if cost is not None else 'NO DETERMINADO'}",
            "**Fuente:** PRESUPUESTO TIARA 2026 — COMBUSTIBLE + SALIDAS BARCO",
        ])

    a, b = results
    lines = [
        "### Análisis financiero — Combustible por salida",
        f"**Período 1:** {label(a[0], a[1])}",
        f"**Gasto de combustible:** {money(a[3]) if a[3] is not None else 'NO DETERMINADO'}",
        f"**Salidas de barco:** {int(a[4]) if a[4] is not None and float(a[4]).is_integer() else a[4] if a[4] is not None else 'NO DETERMINADO'}",
        f"**Combustible promedio por salida:** {money(a[2]) if a[2] is not None else 'NO DETERMINADO'}",
        "",
        f"**Período 2:** {label(b[0], b[1])}",
        f"**Gasto de combustible:** {money(b[3]) if b[3] is not None else 'NO DETERMINADO'}",
        f"**Salidas de barco:** {int(b[4]) if b[4] is not None and float(b[4]).is_integer() else b[4] if b[4] is not None else 'NO DETERMINADO'}",
        f"**Combustible promedio por salida:** {money(b[2]) if b[2] is not None else 'NO DETERMINADO'}",
    ]
    if a[2] is None or b[2] is None:
        lines += ["", "**Diferencia USD por salida:** NO DETERMINADO", "**Variación porcentual:** NO DETERMINADA"]
    else:
        diff = b[2] - a[2]
        pct = (diff / a[2] * 100) if a[2] != 0 else None
        lines += [
            "",
            f"**Diferencia USD por salida:** {money(abs(diff))}",
            f"**Variación porcentual:** {pct:+.2f}%" if pct is not None else
            "**Variación porcentual:** NO DETERMINADA porque el valor de referencia es $0.",
        ]
    lines.append("**Fuente:** PRESUPUESTO TIARA 2026 — COMBUSTIBLE + SALIDAS BARCO")
    return "\n".join(lines)


def _answer_salidas_barco(data, question, thought):
    """Consulta directa de SALIDAS BARCO por mes o año."""
    specs = _period_specs_for_salidas(question, thought)
    def fmt(value):
        if value is None:
            return "NO DETERMINADO"
        try:
            return str(int(value)) if float(value).is_integer() else f"{value:g}"
        except Exception:
            return "NO DETERMINADO"
    if len(specs) >= 2:
        lines = ["### Salidas de barco"]
        for year, month in specs[:2]:
            lines.append(f"**{_financial_period_label(year, month)}:** {fmt(_salidas_barco_value(data, year, month))}")
        lines.append("**Fuente:** PRESUPUESTO TIARA 2026 — SALIDAS BARCO")
        return "\n".join(lines)
    year, month = specs[0]
    return "\n".join([
        "### Salidas de barco",
        f"**Período:** {_financial_period_label(year, month)}",
        f"**Salidas registradas:** {fmt(_salidas_barco_value(data, year, month))}",
        "**Fuente:** PRESUPUESTO TIARA 2026 — SALIDAS BARCO",
    ])


def _financial_analysis(data, thought):
    """Motor del Proyecto de Finanzas: consulta, comparación, promedio y cruces aprobados."""
    p = thought.get("normalized", "")
    rubric = thought.get("rubric") or identify_rubric(p)
    operation = thought.get("financial_operation") or detect_financial_operation(p)

    # SALIDAS BARCO es un rubro de cantidad. Las consultas/comparaciones
    # simples de salidas deben resolverse directamente desde su fila, sin
    # entrar al motor financiero monetario.
    if rubric == "SALIDAS BARCO" and operation == "comparacion":
        return _answer_salidas_barco(data, thought.get("original_question", p), thought)

    # Su único cruce financiero autorizado es combustible ÷ salidas,
    # expresado en USD por salida.
    # La pregunta puede mencionar primero "combustible"; por eso validamos
    # la existencia real del rubro SALIDAS BARCO en el presupuesto antes de
    # activar este cruce, sin depender del orden de las palabras.
    if operation == "costo_combustible_por_salida" and _salidas_barco_record(data) is not None:
        return _answer_combustible_por_salida(data, thought.get("original_question", p), thought)

    if not rubric or not operation or _find_budget_record(data, rubric) is None:
        return None

    month = thought.get("month")
    if operation == "comparacion":
        old_year, new_year = _financial_reference_years(thought)
        # Si hay mes, comparar ese mes. Si la pregunta dice "misma fecha/a hoy",
        # el presupuesto se interpreta al nivel mensual y se muestra además YTD.
        old_value, new_value = _financial_values_for_period(data, rubric, (old_year, new_year), month).get(old_year), _financial_values_for_period(data, rubric, (old_year, new_year), month).get(new_year)
        label = MONTH_NAMES[month].capitalize() if month else "año completo"
        lines = ["### Análisis financiero — Comparación", f"**Rubro:** {rubric}", f"**Período:** {label}",
                 f"**{old_year}:** {money(old_value) if old_value is not None else 'NO DETERMINADO'}",
                 f"**{new_year}:** {money(new_value) if new_value is not None else 'NO DETERMINADO'}"]
        diff, pct = _financial_change(old_value, new_value)
        if diff is None:
            lines += ["**Diferencia:** NO DETERMINADO", "**Variación porcentual:** NO DETERMINADA"]
        else:
            lines.append(f"**Diferencia:** {money(abs(diff))} — el gasto {'aumentó' if diff > 0 else 'disminuyó' if diff < 0 else 'no cambió'}.")
            lines.append(f"**Variación porcentual:** {pct:+.2f}%" if pct is not None else "**Variación porcentual:** NO DETERMINADA porque el valor de referencia es $0.")
        if month is None and any(x in p for x in ("misma fecha", "fecha de hoy", "a hoy", "hasta hoy", "al dia de hoy", "al día de hoy")):
            cm = datetime.now(APP_TIMEZONE).month
            a, na = _financial_ytd(data, rubric, old_year, cm)
            b, nb = _financial_ytd(data, rubric, new_year, cm)
            lines.append(f"**Corte:** hasta {MONTH_NAMES[cm].capitalize()} (el presupuesto se registra por mes).")
            lines.append(f"**Acumulado hasta {MONTH_NAMES[cm].capitalize()}:** {old_year}: {money(a) if a is not None else 'NO DETERMINADO'} · {new_year}: {money(b) if b is not None else 'NO DETERMINADO'}")
            if a is not None and b is not None:
                yd, yp = _financial_change(a, b)
                lines.append(f"**Variación acumulada:** {yp:+.2f}%" if yp is not None else "**Variación acumulada:** NO DETERMINADA")
        lines.append("**Fuente:** PRESUPUESTO TIARA 2026")
        return "\n".join(lines)

    if operation == "promedio":
        year = thought.get("year") or 2026
        monthly = [budget_value_for_month(data, rubric, year, m) for m in range(1, 13)]
        valid = [v for v in monthly if v is not None]
        if not valid:
            return f"### Análisis financiero — Promedio\n**Rubro:** {rubric}\n**Año:** {year}\n**Promedio:** NO DETERMINADO"
        avg = sum(valid) / len(valid)
        return (f"### Análisis financiero — Promedio\n**Rubro:** {rubric}\n**Año:** {year}\n"
                f"**Meses con dato:** {len(valid)} de 12\n**Promedio registrado:** {money(avg)}\n"
                "**Fuente:** PRESUPUESTO TIARA 2026")

    if operation == "gasto_vs_presupuesto":
        year = thought.get("year") or 2026
        budget = budget_value_for_month(data, rubric, year, month) if month else budget_total(data, rubric, year)
        actual = _actual_financial_value(data, rubric, year, month)
        lines = ["### Análisis financiero — Gasto vs presupuesto", f"**Rubro:** {rubric}", f"**Período:** {_financial_period_label(year, month)}",
                 f"**Presupuesto:** {money(budget) if budget is not None else 'NO DETERMINADO'}",
                 f"**Gasto real registrado:** {money(actual) if actual is not None else 'NO DETERMINADO'}"]
        if budget is None or actual is None:
            lines.append("**Diferencia:** NO DETERMINADO — faltan valores comparables.")
            lines.append("**Variación:** NO DETERMINADA")
        else:
            diff = actual - budget
            pct = (diff / budget * 100) if budget != 0 else None
            lines.append(f"**Diferencia gasto − presupuesto:** {money(abs(diff))} — {'por encima' if diff > 0 else 'por debajo' if diff < 0 else 'igual al'} presupuesto.")
            lines.append(f"**Variación:** {pct:+.2f}%" if pct is not None else "**Variación:** NO DETERMINADA porque el presupuesto es $0.")
        lines.append(f"**Fuente de presupuesto:** PRESUPUESTO TIARA 2026\n**Fuente de gasto:** {FINANCIAL_SOURCE_MAP.get(rubric, 'no determinada')}")
        return "\n".join(lines)

    if operation == "consumido_vs_gastado":
        year = thought.get("year") or 2026
        if rubric != "COMSUMO COMBUSTIBLE":
            return (f"### Análisis financiero — Consumido vs gastado\n**Rubro:** {rubric}\n"
                    "Esta operación se puede calcular de forma independiente cuando el Excel contiene cantidad consumida y gasto monetario comparables. Para este rubro no hay dos magnitudes compatibles identificadas.\n"
                    "**Resultado:** NO DETERMINADO")
        records = _source_records(data, "CONSUMO COMBUSTIBLE ")
        liters = 0.0; spent = 0.0; found_l = found_s = False
        for rec in records:
            d = _record_data(rec); dt = _date_from_any(d.get("MES"))
            if not dt or dt.year != year or (month and dt.month != month): continue
            l = _budget_number(d.get("GASOLINA LITROS")); s = _budget_number(d.get("TOTAL"))
            if l is not None: liters += l; found_l = True
            if s is not None: spent += s; found_s = True
        lines = ["### Análisis financiero — Consumido vs gastado", f"**Rubro:** {rubric}", f"**Período:** {_financial_period_label(year, month)}",
                 f"**Combustible consumido:** {f'{liters:,.2f} litros' if found_l else 'NO DETERMINADO'}",
                 f"**Gasto registrado:** {money(spent) if found_s else 'NO DETERMINADO'}"]
        if found_l and found_s and liters > 0:
            lines.append(f"**Costo por litro:** {money(spent / liters)}")
        else:
            lines.append("**Costo por litro:** NO DETERMINADO")
        lines.append("**Fuente:** CONSUMO COMBUSTIBLE")
        return "\n".join(lines)
    return None


def answer_expenses(data, thought):
    p = thought["normalized"]
    rubric = identify_rubric(p)
    year = thought["year"] or 2026
    month = thought["month"]

    # Pregunta mensual sin rubro: responder con el TOTAL GENERAL del mes.
    # Ejemplo: '¿Cuánto se gastó en agosto del 2026?'
    if month and not rubric:
        value = budget_total_all(data, year, month)
        if value is None:
            return (f"### Gasto mensual\n**Mes:** {MONTH_NAMES[month].capitalize()} {year}\n"
                    "**Gasto registrado:** NO DETERMINADO")
        return (f"### Gasto mensual — Presupuesto Tiara\n"
                f"**Mes:** {MONTH_NAMES[month].capitalize()} {year}\n"
                f"**Gasto total registrado:** {money(value)}\n"
                "**Fuente:** PRESUPUESTO TIARA 2026\n"
                "**Estado:** CONFIRMADO")

    # Pregunta anual sin rubro: responder con el total general anual.
    if not month and not rubric:
        value = budget_total_all(data, year)
        if value is None:
            return f"### Gasto anual\n**Año:** {year}\n**Gasto total:** NO DETERMINADO"
        return (f"### Gasto anual — Presupuesto Tiara\n**Año:** {year}\n"
                f"**Gasto total registrado:** {money(value)}\n"
                "**Fuente:** PRESUPUESTO TIARA 2026\n**Estado:** CONFIRMADO")

    if not rubric:
        return "No pude determinar el rubro del gasto. Puedes indicar el concepto (combustible, motores, generador, limpieza, etc.)."

    if month:
        value = budget_value_for_month(data, rubric, year, month)
        if value is None:
            return (f"### Gasto mensual\n**Rubro:** {rubric}\n"
                    f"**Mes:** {MONTH_NAMES[month].capitalize()} {year}\n"
                    "**Gasto registrado:** NO DETERMINADO")
        return (f"### Gasto mensual\n**Rubro:** {rubric}\n"
                f"**Mes:** {MONTH_NAMES[month].capitalize()} {year}\n"
                f"**Gasto registrado:** {money(value)}\n"
                "**Fuente:** PRESUPUESTO TIARA 2026\n**Estado:** CONFIRMADO")

    value = budget_total(data, rubric, year)
    if value is None:
        return f"### Gasto anual\n**Rubro:** {rubric}\n**Año:** {year}\n**Gasto:** NO DETERMINADO"
    return (f"### Gasto anual\n**Rubro:** {rubric}\n**Año:** {year}\n"
            f"**Gasto registrado:** {money(value)}\n"
            "**Fuente:** PRESUPUESTO TIARA 2026\n**Estado:** CONFIRMADO")

def answer_fuel(data, thought):
    year = thought["year"] or 2026
    month = thought.get("month")

    # Cuando la consulta es mensual y viene dentro de una conversación de
    # gastos (por ejemplo: "En ese mismo mes cuanto en combustible"),
    # consultar primero el rubro financiero de combustible del PRESUPUESTO.
    # Esto evita devolver los 12 meses cuando el usuario pidió un solo mes.
    if month:
        # El Excel contiene el rubro escrito de dos formas:
        # "COMSUMO COMBUSTIBLE" y "CONSUMO COMBUSTIBLE". Buscar ambas
        # variantes para no perder el mes solicitado.
        rec = None
        value = None
        for rubric in ("COMSUMO COMBUSTIBLE", "CONSUMO COMBUSTIBLE"):
            rec, value = _find_budget_record_for_month(data, rubric, year, month)
            if rec is not None and value is not None:
                break
        if rec is not None and value is not None:
            return (
                f"### Gasto de combustible\n"
                f"**Mes:** {MONTH_NAMES[month].capitalize()} {year}\n"
                f"**Combustible registrado:** {money(value)}\n"
                f"**Fuente:** PRESUPUESTO TIARA 2026\n"
                f"**Estado:** CONFIRMADO"
            )

    # Si no hay mes, o el presupuesto no contiene el mes solicitado, usar la
    # bitácora de consumo de combustible como respaldo.
    rows = data.get("fuel_records", [])
    totals = {m: 0.0 for m in range(1, 13)}
    found = {m: False for m in range(1, 13)}

    for rec in rows:
        vals = rec.get("values", [])
        if not vals:
            continue
        # En la base normalizada de combustible, MES es la primera columna y
        # contiene el mes como fecha serial de Excel.
        d = parse_date(vals[0]) or excel_serial_to_date(vals[0])
        if not d or d.year != year:
            continue
        total = safe_float(vals[5]) if len(vals) > 5 else None
        if total is not None:
            totals[d.month] += total
            found[d.month] = True

    if month:
        if found[month]:
            return (
                f"### Consumo de combustible\n"
                f"**Mes:** {MONTH_NAMES[month].capitalize()} {year}\n"
                f"**Combustible registrado:** {money(totals[month])}\n"
                f"**Fuente:** CONSUMO COMBUSTIBLE\n"
                f"**Estado:** CONFIRMADO"
            )
        return (
            f"### Consumo de combustible\n"
            f"**Mes:** {MONTH_NAMES[month].capitalize()} {year}\n"
            f"**Combustible registrado:** NO DETERMINADO\n"
            f"**Fuente:** CONSUMO COMBUSTIBLE / PRESUPUESTO TIARA 2026"
        )

    annual = sum(totals.values())
    lines = [f"### Consumo de combustible — {year}"]
    for m in range(1, 13):
        lines.append(
            f"- **{MONTH_NAMES[m].capitalize()}:** "
            f"{money(totals[m]) if found[m] else 'NO DETERMINADO'}"
        )
    lines.append(f"\n**Total del año:** {money(annual)}")
    return "\n".join(lines)


# ============================================================
# INVENTARIO / BITÁCORA / DOCUMENTOS
# ============================================================

def answer_inventory(data, thought):
    p = thought["normalized"]
    rows = data.get("inventory_records", [])
    if not rows:
        return "No hay datos de inventario disponibles."
    tokens = [t for t in re.findall(r"[a-z0-9]+", p) if len(t) >= 4 and t not in STOPWORDS]
    hits = []
    for rec in rows:
        vals = rec.get("values", [])
        if len(vals) < 3:
            continue
        product = norm(vals[1])
        score = sum(1 for t in tokens if t in product)
        if score:
            hits.append((score, product, vals))
    hits.sort(reverse=True)
    if not hits:
        return "No encontré ese producto en el inventario."
    lines = ["### Inventario"]
    for _, _, vals in hits[:8]:
        lines.append(f"- **{vals[1]}** — existencias: **{vals[2]}**")
    return "\n".join(lines)


def answer_log(data, thought):
    p = thought["normalized"]
    rows = log_records(data)
    if not rows:
        return "No hay registros de bitácora disponibles."

    # Filtros de estado: primero usamos la estructura de la bitácora y solo
    # después el texto libre.
    wants_pending = any(x in p for x in ("pendiente", "pendientes", "sin resolver", "no resuelto"))
    wants_done = any(x in p for x in ("resuelto", "resueltos", "listo", "reparado", "reparados", "atendido"))
    system = thought.get("system")
    if system not in {"Motores Mercury", "Patas Mercury", "Generador Cummins Onan"}:
        system = None

    candidates = []
    for rec in rows:
        d = rec.get("data") or {}
        if system and d.get("sistema") != system:
            continue
        if wants_pending and d.get("estado") != "PENDIENTE":
            continue
        if wants_done and d.get("estado") != "RESUELTO":
            continue
        candidates.append(rec)

    # Si preguntó explícitamente por pendientes/resueltos, ordenar por fecha
    # y mostrar primero los eventos más recientes.
    if wants_pending or wants_done:
        candidates.sort(key=lambda r: _log_date(r) or date.min, reverse=True)
        if not candidates:
            estado = "pendientes" if wants_pending else "resueltos"
            return f"No encontré eventos {estado} en la bitácora para ese sistema."
        title = "Pendientes" if wants_pending else "Resueltos / atendidos"
        lines = [f"### Bitácora — {title}"]
        for rec in candidates[:15]:
            lines.append(f"- {_format_log_record(rec)}")
        return "\n".join(lines)

    # Consulta normal: búsqueda relevante, pero sin devolver filas ajenas solo
    # porque compartan palabras genéricas como 'motor' o 'generador'.
    tokens = [t for t in re.findall(r"[a-z0-9]+", p) if len(t) >= 4 and t not in STOPWORDS]
    hits = []
    for rec in candidates:
        d = rec.get("data") or {}
        text = norm(" ".join(str(d.get(k) or "") for k in ("observacion", "observaciones", "sistema", "estado")))
        score = 0
        for t in tokens:
            if t in text:
                score += 1
        if score:
            hits.append((score, _log_date(rec) or date.min, rec))
    hits.sort(key=lambda x: (x[0], x[1]), reverse=True)
    if not hits:
        return "No encontré una entrada de bitácora relacionada con la consulta."

    lines = ["### Bitácora — resultados relevantes"]
    for _, _, rec in hits[:8]:
        lines.append(f"- {_format_log_record(rec)}")
    return "\n".join(lines)


def answer_documents(data):
    records = data.get("permit_records") or []
    if records:
        lines = ["### Permisos y seguros"]
        for r in records:
            raw = r.get("fecha_vencimiento")
            if raw:
                try:
                    fecha = fmt_date(datetime.fromisoformat(raw).date())
                except Exception:
                    fecha = raw
                lines.append(f"- **{r['documento']}** — vence **{fecha}**.")
            else:
                lines.append(f"- **{r['documento']}** — fecha de vencimiento no registrada en las celdas.")
        return "\n".join(lines)
    rows = data.get("documents_records", [])
    if not rows:
        return "No hay información de permisos y seguros disponible."
    lines = ["### Permisos y seguros"]
    for rec in rows:
        vals = rec.get("values", [])
        text = " | ".join(str(v) for v in vals if v not in (None, ""))
        if text:
            lines.append("- " + text)
    return "\n".join(lines[:15])


def answer_permit_expiry(data, question):
    """Determina el próximo vencimiento desde las celdas del Excel primero."""
    import datetime as _dt
    records = data.get("permit_records") or []
    today = _dt.date.today()
    structured = []
    for r in records:
        raw = r.get("fecha_vencimiento")
        if not raw:
            continue
        try:
            dt = _dt.date.fromisoformat(raw)
        except Exception:
            continue
        structured.append((dt, r))
    if not structured:
        return answer_photo_expiry(data, question)
    future = sorted([x for x in structured if x[0] >= today], key=lambda x: x[0])
    candidates = future if future else sorted(structured, key=lambda x: x[0])
    dt, rec = candidates[0]
    lines = [
        "### Próximo vencimiento — Permisos & Seguros",
        f"**Documento:** {rec['documento']}",
        f"**Fecha de vencimiento:** {fmt_date(dt)}",
        f"**Celda del Excel:** {rec.get('celda_vencimiento') or 'no determinada'}",
        "\n**Vencimientos registrados:**",
    ]
    for odt, orc in sorted(structured, key=lambda x: x[0]):
        mark = " ← **más próximo**" if orc is rec else ""
        lines.append(f"- {orc['documento']}: {fmt_date(odt)}{mark}")
    missing = [r["documento"] for r in records if not r.get("fecha_vencimiento")]
    if missing:
        lines.append("\n**Sin fecha escrita en la hoja:** " + ", ".join(missing) + ".")
    return "\n".join(lines)


def answer_sheet(data, sheet_name, title):
    rows = sheet_rows(data, sheet_name)
    if not rows:
        return f"No hay información disponible en la hoja **{sheet_name}**."
    lines = [f"### {title}"]
    for row in rows[:20]:
        vals = row.get("values", []) if isinstance(row, dict) else row
        text = " | ".join(str(v) for v in vals if v not in (None, ""))
        if text: lines.append("- " + text)
    return "\n".join(lines)


def _parse_vision_date(raw):
    """Convierte fechas comunes devueltas por el análisis visual."""
    raw = str(raw or "").strip().strip(".,;")
    if not raw:
        return None
    try:
        if re.fullmatch(r"\d{4}-\d{2}(?:-\d{2})?", raw):
            return _dt.date.fromisoformat(raw if len(raw) == 10 else raw + "-01")
        m = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", raw)
        if m:
            a, b, y = map(int, m.groups())
            if y < 100:
                y += 2000
            # Los registros de Florida/EE. UU. usan normalmente MM/DD/YYYY.
            if a <= 12 and b <= 31:
                return _dt.date(y, a, b)
    except Exception:
        return None
    return None


def _vision_expiry_dates(data):
    """Extrae vencimientos del formato nuevo y de análisis visuales antiguos.

    Así una base ya guardada sigue siendo útil aunque Gemini hubiera devuelto
    antes un texto narrativo en lugar de FECHA_VENCIMIENTO: YYYY-MM-DD.
    """
    import datetime as _dt
    found = []
    label = (
        r"(?:FECHA[_ ]?VENCIMIENTO|FECHA DE VENCIMIENTO|VENCIMIENTO|"
        r"EXPIR(?:ES|ATION|Y)|EXPIRATION DATE|EXPIRES|VALID UNTIL|"
        r"COVERAGE ENDS|COVERAGE END|POLICY ENDS|VIGENCIA HASTA|"
        r"FECHA FIN|ENDS|RENEWAL)"
    )
    date_token = r"(\d{4}-\d{2}(?:-\d{2})?|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})"
    for im in data.get("images", []):
        if norm(im.get("sheet", "")) != "permisos & seguros":
            continue
        analysis = str(im.get("vision_analysis") or "")
        if not analysis:
            continue
        candidates = []
        # Formato estructurado de las versiones nuevas.
        for m in re.finditer(r"FECHA[_ ]?VENCIMIENTO\s*:\s*" + date_token, analysis, re.I):
            candidates.append(m.group(1))
        # Texto narrativo de versiones anteriores.
        if not candidates:
            for m in re.finditer(label + r"[^\n]{0,70}?" + date_token, analysis, re.I):
                candidates.append(m.group(1))
        for raw in candidates:
            dt = _parse_vision_date(raw)
            if dt:
                found.append((dt, im, raw))
                break
    return found

def answer_photo_expiry(data, question):
    found = _vision_expiry_dates(data)
    if not found:
        return (
            "### Vencimientos en fotografías\n"
            "No encontré una fecha de vencimiento explícitamente legible en los "
            "análisis visuales guardados. Si las fotos fueron analizadas con una "
            "versión anterior del agente, hay que procesarlas una sola vez con la "
            "lectura enfocada en fechas; después quedarán guardadas para futuras "
            "consultas."
        )
    import datetime as _dt
    today = _dt.date.today()
    future = sorted([x for x in found if x[0] >= today], key=lambda x: x[0])
    candidates = future if future else sorted(found, key=lambda x: x[0])
    dt, im, raw = candidates[0]
    return (
        "### Próximo vencimiento\n"
        f"- **Fecha:** {raw}\n"
        f"- **Fotografía:** {im.get('index','?')} — hoja **{im.get('sheet','PERMISOS & SEGUROS')}**, fila {im.get('row','?')}\n"
        f"- **Análisis visual:** {str(im.get('vision_analysis') or '').replace(chr(10), ' ')[:500]}"
    )

def answer_photos(data, question):
    p = norm(question)
    if any(x in p for x in (
        "vencimiento mas proximo", "vencimiento más próximo", "vence primero",
        "vence antes", "proximo vencimiento", "próximo vencimiento",
        "fecha de vencimiento mas proxima", "fecha de vencimiento más próxima",
        "cual vence", "cuál vence"
    )):
        return answer_photo_expiry(data, question)
    images = data.get("images", [])
    if not images:
        return "No encontré fotografías incrustadas en el Excel."
    target = norm(question)
    preferred = [im for im in images if norm(im.get("sheet", "")) in target or ("permiso" in target and "permiso" in norm(im.get("sheet", "")))]
    chosen = preferred or images
    lines = [f"### Fotografías — {len(images)} encontradas"]
    for im in chosen[:12]:
        loc = f"hoja **{im.get('sheet','desconocida')}**"
        if im.get("row"): loc += f", fila {im['row']}"
        lines.append(f"\n**Fotografía {im.get('index','?')}** — {loc}")
        analysis = str(im.get("vision_analysis") or "").strip()
        if analysis and not analysis.startswith("No se pudo"):
            lines.append(analysis)
        else:
            err = str(im.get("vision_error") or "").strip()
            if err:
                lines.append(f"⚠️ No se pudo interpretar esta fotografía: {err}")
            else:
                lines.append("No hay interpretación visual guardada para esta fotografía.")
    vision = data.get("metadata", {}).get("vision", {})
    if not vision.get("enabled"):
        lines.append("\n⚠️ La lectura visual todavía no está conectada. Configura el secret `gemini.api_key` en Streamlit para activar la interpretación de las fotos.")
    return "\n".join(lines)


def answer_general(data, thought):
    if thought["system"] == "Combustible":
        return answer_fuel(data, thought)
    if thought["system"] == "Inventario": return answer_inventory(data, thought)
    if thought["system"] == "Facturas": return answer_sheet(data, "FACTURAS 2026...", "Facturas")
    if thought["system"] == "Limpieza": return answer_sheet(data, "S.LIMPIEZA & OTROS", "Limpieza y otros")
    if thought["system"] == "Permisos y seguros": return answer_documents(data)
    return answer_sheet(data, thought.get("system") or "", "Información") if thought.get("system") not in (None, "General") else ("Puedo consultar las 8 hojas del Excel: mantenimiento, bitácora, inventario, combustible, presupuesto, facturas, limpieza y permisos/seguros.\n\nPrueba: **¿Cuánto gastamos en limpieza en septiembre?**")


def _is_financial_capabilities_question(question):
    """Reconoce la consulta especial que pide conocer las capacidades financieras."""
    p = norm(question)
    # No depende del Excel ni de un rubro: es una pregunta fija de capacidades.
    return bool(
        re.search(r"\bcapacidades?\b", p)
        and re.search(r"\banalisis financiero\b", p)
    ) or p in {
        "que puedes hacer en analisis financiero",
        "que puedes hacer con el analisis financiero",
        "cuales son tus capacidades financieras",
        "cuales son tus capacidades de analisis financiero",
    }


def _financial_capabilities_answer():
    """Respuesta breve y estable para explicar las capacidades financieras."""
    return (
        "### Capacidades de análisis financiero\n"
        "Puedo realizar: comparación de gastos entre años o meses; diferencias en dólares; "
        "aumentos o disminuciones porcentuales; promedios; comparación de gasto contra "
        "presupuesto; comparación de consumido contra gastado cuando existan ambos datos; "
        "y análisis de un rubro cruzando el presupuesto con su información correspondiente.\n\n"
        "Si falta un dato necesario, lo indicaré como **NO DETERMINADO** y no inventaré el resultado."
    )


def run_agent(data, question, context=None):
    if _is_financial_capabilities_question(question):
        return _financial_capabilities_answer()

    thought = think(question)

    # Contexto conversacional INVISIBLE: solo conserva datos necesarios para
    # interpretar referencias como "ese mismo mes". No guarda ni muestra
    # preguntas ni respuestas anteriores.
    context = context or {}
    p = thought["normalized"]
    refers_to_previous_month = bool(re.search(
        r"\b(ese mismo mes|ese mes|en ese mismo mes|en ese mes|mismo mes)\b", p
    ))
    if refers_to_previous_month and thought.get("month") is None:
        thought["month"] = context.get("month")
        if thought.get("year") is None:
            thought["year"] = context.get("year")

    # SALIDAS BARCO se resuelve directamente desde PRESUPUESTO TIARA 2026.
    # No pasa por rutas genéricas ni por servicios externos.
    if thought.get("rubric") == "SALIDAS BARCO":
        op = thought.get("financial_operation")
        if op == "costo_combustible_por_salida":
            financial = _financial_analysis(data, thought)
            if financial is not None:
                return financial
        # Las consultas directas de SALIDAS BARCO (mes, año y comparación)
        # deben responderse desde la fila SALIDAS BARCO de PRESUPUESTO TIARA 2026.
        return _answer_salidas_barco(data, question, thought)

    # Proyecto de Finanzas para los demás rubros.
    financial = _financial_analysis(data, thought)
    if financial is not None:
        return financial

    intent = thought["intent"]

    if intent == "estado_general":
        return answer_state(data)
    if intent == "ultimo_mantenimiento":
        return answer_ultimo_mantenimiento(data, thought)
    if intent == "historial_completo":
        return answer_historial(data, thought, complete=True)
    if intent == "cantidad_mantenimiento":
        return answer_count(data, thought)
    if intent == "kit_300":
        return answer_kit(data)
    if intent == "mantenimiento":
        return answer_mantenimiento(data, thought)
    if intent in ("gasto_mensual", "gasto_anual"):
        return answer_expenses(data, thought)
    if intent == "combustible":
        return answer_fuel(data, thought)
    if intent == "inventario":
        return answer_inventory(data, thought)
    if intent == "bitacora":
        return answer_log(data, thought)
    if intent == "documentos":
        if any(x in p for x in (
            "vencimiento", "vence", "expira", "expiracion", "expiración",
            "fecha limite", "fecha límite", "vigencia hasta", "vence primero",
            "mas proximo", "más próximo", "fecha de vencimiento"
        )):
            return answer_permit_expiry(data, question)
        if any(x in p for x in ("foto", "fotografia", "fotografía", "imagen", "que dice", "qué dice", "lee", "leer", "muestra")):
            return answer_photos(data, question)
        return answer_documents(data)
    if intent == "fotografias":
        if any(x in p for x in ("vencimiento", "vence", "vence primero", "mas proximo", "más próximo", "fecha de vencimiento")):
            return answer_permit_expiry(data, question)
        return answer_photos(data, question)
    if intent == "facturas":
        return answer_sheet(data, "FACTURAS 2026...", "Facturas")
    if intent == "limpieza":
        return answer_sheet(data, "S.LIMPIEZA & OTROS", "Limpieza y otros")
    if intent == "fotografias":
        return answer_photos(data, question)
    if intent == "horometros":
        h = current_hours(data)
        return "### Horómetros actuales\n" + "\n".join(f"- **{k}:** {v} h" if v is not None else f"- **{k}:** NO DETERMINADO" for k, v in h.items())
    return answer_general(data, thought)


# ============================================================
# INTERFAZ STREAMLIT — OSCURA Y SIN HISTORIAL
# ============================================================

st.set_page_config(
    page_title="Agente Tiara",
    page_icon="⚓",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ============================================================
# INTERFAZ PRINCIPAL
# ============================================================
# Regla de diseño:
# - El sidebar sigue siendo 100 % nativo de Streamlit.
# - Solo se modifica la presentación de la pantalla principal.
# - Los cinco accesos rápidos siguen entrando en run_agent().
# - Texto y voz siguen usando exactamente el mismo cerebro.
st.markdown("""
<style>
html, body, [data-testid="stAppViewContainer"], [data-testid="stApp"] {
    background: #07111d !important;
}

[data-testid="stHeader"] {
    background: #07111d !important;
}

.block-container {
    max-width: 1050px !important;
    padding-top: 3.5rem !important;
    padding-bottom: 2.2rem !important;
}

/* Encabezado */
.tiara-header {
    display: flex;
    align-items: center;
    justify-content: center;
    gap: 0;
    margin: 0 auto 12px auto;
    padding: 6px 0 0 0;
}

.tiara-boat {
    font-size: 3.15rem;
    line-height: 1;
    filter: drop-shadow(0 3px 4px rgba(0,0,0,.35));
}

.tiara-brand {
    text-align: left;
    line-height: 1;
}

.tiara-brand-title {
    color: #f4f7fb;
    font-size: clamp(1.25rem, 2.6vw, 1.75rem);
    font-weight: 850;
    letter-spacing: .03em;
    white-space: nowrap;
}

.tiara-brand-title span {
    color: #19b5f1;
}

.tiara-brand-subtitle {
    margin-top: 9px;
    color: #b9cbd8;
    font-size: clamp(.72rem, 1.8vw, .92rem);
    font-weight: 700;
    letter-spacing: .22em;
}

/* Presentación / hero */
.tiara-hero {
    background: linear-gradient(135deg, #10446a 0%, #09689c 100%);
    border: 1px solid rgba(42, 178, 239, .35);
    border-radius: 25px;
    padding: 6px 20px;
    margin: 0 0 18px 0;
    box-shadow: inset 0 1px 0 rgba(255,255,255,.05);
}

.tiara-hero p {
    color: #e7f2fa !important;
    font-size: clamp(.88rem, 1.7vw, 1.08rem) !important;
    line-height: 1.28 !important;
    margin: 0 !important;
}

/* Título de consultas */
.tiara-section-title {
    color: #f0f5f9;
    font-size: clamp(1.35rem, 3vw, 1.75rem);
    font-weight: 800;
    margin: 0 0 15px 7px;
}

/* Botones rápidos: el layout 2 + 2 + 1 se consigue con dos filas
   de columnas y un botón independiente de ancho completo. */
.quick-row {
    margin-bottom: 12px;
}

div.stButton > button {
    min-height: 74px !important;
    height: 74px !important;
    border-radius: 22px !important;
    border: 1px solid #36536b !important;
    background: linear-gradient(145deg, #111c28 0%, #0d1722 100%) !important;
    color: #e9f0f5 !important;
    font-size: clamp(1rem, 2.5vw, 1.25rem) !important;
    font-weight: 550 !important;
    letter-spacing: .01em !important;
    box-shadow: inset 0 1px 0 rgba(255,255,255,.025) !important;
    transition: border-color .15s ease, transform .15s ease, box-shadow .15s ease !important;
}

div.stButton > button:hover {
    border-color: #159ddd !important;
    box-shadow: 0 0 0 1px rgba(21,157,221,.14), 0 8px 22px rgba(0,0,0,.18) !important;
    transform: translateY(-1px);
}

div.stButton > button:active {
    transform: translateY(0);
}

/* Zona de chat */
.tiara-chat-space {
    height: 15px;
}

[data-testid="stChatMessage"] {
    border-radius: 22px !important;
    border: 1px solid #29445b !important;
    background: linear-gradient(145deg, #121d29 0%, #0d1722 100%) !important;
    margin: 12px 0 !important;
    padding: 16px 18px !important;
}

[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] p,
[data-testid="stChatMessage"] [data-testid="stMarkdownContainer"] li {
    color: #eaf1f6 !important;
}

/* Burbuja de usuario ligeramente azul, como en la referencia. */
[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) {
    background: linear-gradient(135deg, #075bb8 0%, #087bd8 100%) !important;
    border-color: #0a94e7 !important;
}

/* Campo de respuesta del agente */
.tiara-response-label {
    color: #aec3d2;
    font-size: .82rem;
    font-weight: 700;
    letter-spacing: .08em;
    margin: 18px 0 8px 6px;
}

.st-key-tiara_response {
    border: 1px solid #36536b !important;
    background: linear-gradient(145deg, #111c28 0%, #0d1722 100%) !important;
    color: #eaf1f6 !important;
    border-radius: 20px !important;
    padding: 18px 20px !important;
    min-height: 70px;
    line-height: 1.55;
    font-size: 1rem;
    box-shadow: inset 0 1px 0 rgba(255,255,255,.025) !important;
    overflow-wrap: anywhere;
}

.st-key-tiara_response [data-testid="stMarkdownContainer"] {
    color: #eaf1f6 !important;
}

.st-key-tiara_response [data-testid="stMarkdownContainer"] p,
.st-key-tiara_response [data-testid="stMarkdownContainer"] li {
    color: #eaf1f6 !important;
}

.st-key-tiara_response [data-testid="stMarkdownContainer"] strong {
    color: #ffffff !important;
}

/* Compositor de texto/voz */
/* El compositor ya no depende de st.form para poder procesar el audio
   de forma fiable, pero conserva visualmente la misma caja. */
.st-key-tiara_composer {
    position: fixed !important;
    left: 50% !important;
    bottom: 10px !important;
    transform: translateX(-50%) !important;
    width: min(700px, calc(100vw - 24px)) !important;
    max-width: calc(100vw - 24px) !important;
    z-index: 9999 !important;
    background: rgba(15, 29, 43, .96) !important;
    backdrop-filter: blur(14px);
    -webkit-backdrop-filter: blur(14px);
    border: 1px solid #35536a !important;
    border-radius: 24px !important;
    padding: 8px 10px 7px !important;
    margin: 0 !important;
    box-shadow: 0 10px 35px rgba(0,0,0,.38) !important;
    box-sizing: border-box !important;
}

.st-key-tiara_composer [data-testid="stHorizontalBlock"] {
    flex-wrap: nowrap !important;
    align-items: center !important;
    gap: 8px !important;
}

.st-key-tiara_composer [data-testid="stHorizontalBlock"] > div:nth-child(1),
.st-key-tiara_composer [data-testid="stHorizontalBlock"] > div:nth-child(3) {
    flex: 0 0 52px !important;
    width: 52px !important;
    min-width: 52px !important;
}

.st-key-tiara_composer [data-testid="stHorizontalBlock"] > div:nth-child(2) {
    flex: 1 1 auto !important;
    width: auto !important;
    min-width: 0 !important;
}

.st-key-tiara_composer div[data-testid="stAudioInput"],
.st-key-tiara_composer div[data-testid="stTextInput"],
.st-key-tiara_composer .st-key-tiara_send_question {
    width: 100% !important;
    min-width: 0 !important;
}

.tiara-composer-label {
    color: #aec3d2;
    font-size: .95rem;
    margin: 8px 0 7px 8px;
}

[data-testid="stForm"] {
    background: rgba(15, 29, 43, .78) !important;
    border: 1px solid #35536a !important;
    border-radius: 27px !important;
    padding: 9px 11px !important;
    margin-top: 8px !important;
}

/* Campo de texto */
div[data-testid="stTextInput"] input {
    min-height: 58px !important;
    border-radius: 30px !important;
    background: #172638 !important;
    border: 1px solid #314d65 !important;
    color: #f2f6f9 !important;
    font-size: 1.03rem !important;
    padding-left: 20px !important;
}

div[data-testid="stTextInput"] input::placeholder {
    color: #9eb1c0 !important;
    opacity: 1 !important;
}

/* Micrófono: en versiones antiguas usamos st.audio_input como fallback.
   El reproductor que Streamlit muestra después de grabar no forma parte del
   compositor: ocultamos solo esa previsualización, sin ocultar el botón ni
   afectar el UploadedFile que recibe Python. */
div[data-testid="stAudioInput"] {
    min-height: 58px !important;
    height: 58px !important;
    max-height: 58px !important;
    overflow: hidden !important;
}

div[data-testid="stAudioInput"] button {
    min-height: 58px !important;
    height: 58px !important;
    width: 100% !important;
    border-radius: 30px !important;
    background: #182b3e !important;
    border: 1px solid #607a90 !important;
}

div[data-testid="stAudioInput"] audio,
div[data-testid="stAudioInput"] [data-testid="stAudio"] {
    display: none !important;
}

/* Chat nativo de Streamlit: es el compositor principal cuando la versión
   instalada soporta audio integrado. Queda anclado abajo por Streamlit. */
div[data-testid="stChatInput"] {
    width: min(700px, calc(100vw - 24px)) !important;
    max-width: calc(100vw - 24px) !important;
    margin-left: auto !important;
    margin-right: auto !important;
}

div[data-testid="stChatInput"] > div {
    background: rgba(15, 29, 43, .96) !important;
    border: 1px solid #35536a !important;
    border-radius: 24px !important;
    box-shadow: 0 10px 35px rgba(0,0,0,.38) !important;
    backdrop-filter: blur(14px);
    -webkit-backdrop-filter: blur(14px);
}

div[data-testid="stChatInputTextArea"] {
    color: #f2f6f9 !important;
}

div[data-testid="stChatInputTextArea"]::placeholder {
    color: #9eb1c0 !important;
    opacity: 1 !important;
}

div[data-testid="stChatInput"] button {
    border-radius: 50% !important;
}

/* El audio integrado del chat mantiene el botón pequeño, igual que enviar. */
div[data-testid="stChatInput"] [data-testid*="Audio"] button,
div[data-testid="stChatInput"] button[aria-label*="audio" i],
div[data-testid="stChatInput"] button[aria-label*="record" i],
div[data-testid="stChatInput"] button[aria-label*="micro" i] {
    width: 40px !important;
    height: 40px !important;
    min-width: 40px !important;
    min-height: 40px !important;
}

div[data-testid="stChatInput"] [data-testid="stChatInputSubmitButton"] button,
div[data-testid="stChatInput"] button[data-testid="stChatInputSubmitButton"] {
    width: 40px !important;
    height: 40px !important;
    min-width: 40px !important;
    min-height: 40px !important;
}

/* Botón enviar */
[data-testid="stFormSubmitButton"] button {
    min-height: 58px !important;
    height: 58px !important;
    border-radius: 30px !important;
    background: linear-gradient(135deg, #078ce8 0%, #1475e4 100%) !important;
    border: 1px solid #4db8ff !important;
    color: white !important;
    font-size: 1.65rem !important;
    font-weight: 800 !important;
    padding: 0 !important;
    box-shadow: 0 5px 20px rgba(0,130,235,.22) !important;
}

[data-testid="stFormSubmitButton"] button:hover {
    background: linear-gradient(135deg, #129cf0 0%, #1681ef 100%) !important;
    border-color: #80ceff !important;
}

/* Boton enviar cuando el compositor usa st.button en vez de st.form. */
.st-key-tiara_send_question button {
    min-height: 58px !important;
    height: 58px !important;
    border-radius: 30px !important;
    background: linear-gradient(135deg, #078ce8 0%, #1475e4 100%) !important;
    border: 1px solid #4db8ff !important;
    color: white !important;
    font-size: 1.65rem !important;
    font-weight: 800 !important;
    padding: 0 !important;
    box-shadow: 0 5px 20px rgba(0,130,235,.22) !important;
}

.st-key-tiara_send_question button:hover {
    background: linear-gradient(135deg, #129cf0 0%, #1681ef 100%) !important;
    border-color: #80ceff !important;
}

.tiara-help {
    color: #a8bdcd !important;
    font-size: .96rem !important;
    line-height: 1.55 !important;
    margin: 8px 8px 0 8px !important;
}

[data-testid="stMarkdownContainer"] p,
[data-testid="stMarkdownContainer"] li {
    color: #e7edf2;
}

small, .stCaption {
    color: #9da9b5 !important;
}

/* En móvil: mantener el mismo aspecto y evitar que Streamlit comprima
   demasiado los controles. */
@media (max-width: 700px) {
    .block-container {
        padding-left: 18px !important;
        padding-right: 18px !important;
        padding-bottom: 125px !important;
        overflow-x: hidden !important;
        padding-top: 3.2rem !important;
    }

    /* Encabezado móvil: en pantallas de ~6.1 pulgadas se apila para que
       AGENTE TIARA siempre quede completamente visible, sin comprimir ni
       recortar el resto de la interfaz. */
    .tiara-header {
        flex-direction: column;
        justify-content: center;
        align-items: center;
        gap: 4px;
        margin-bottom: 12px;
        padding: 0 4px;
        width: 100%;
        box-sizing: border-box;
    }

    .tiara-brand {
        width: 100%;
        text-align: center;
    }

    .tiara-brand-title {
        font-size: clamp(1.0rem, 4.8vw, 1.28rem);
        line-height: 1.08;
        white-space: nowrap;
    }

    .tiara-brand-subtitle {
        font-size: .55rem;
        letter-spacing: .14em;
        margin-top: 4px;
        white-space: nowrap;
    }

    .tiara-hero {
        border-radius: 20px;
        padding: 6px 14px;
        margin-bottom: 12px;
    }

    .tiara-hero p {
        font-size: .78rem !important;
        line-height: 1.25 !important;
    }

    .tiara-section-title {
        font-size: 1.34rem;
        margin-left: 5px;
    }

    div.stButton > button {
        min-height: 72px !important;
        height: 72px !important;
        border-radius: 21px !important;
        font-size: 1rem !important;
    }

    [data-testid="stForm"] {
        border-radius: 25px !important;
        padding: 8px !important;
    }

    div[data-testid="stTextInput"] input,
    div[data-testid="stAudioInput"] button,
    [data-testid="stFormSubmitButton"] button {
        min-height: 58px !important;
        height: 58px !important;
    }

    .tiara-help {
        font-size: .9rem !important;
    }
}
</style>
""", unsafe_allow_html=True)

# ============================================================
# CARGA DEL EXCEL DESDE LA PROPIA APLICACIÓN
# ============================================================
# El usuario puede reemplazar el Excel sin tocar GitHub.
# El .xlsx/.xlsm se procesa internamente: no hay que descomprimirlo.
if "uploaded_data" not in st.session_state:
    st.session_state.uploaded_data = None
if "uploaded_signature" not in st.session_state:
    st.session_state.uploaded_signature = None
if "upload_message" not in st.session_state:
    st.session_state.upload_message = ""

with st.sidebar:
    st.markdown("### 📥 Actualizar Excel")
    st.caption(
        "Carga el Excel nuevo del barco. La aplicación lo procesa internamente "
        "y guarda la base activa para futuras sesiones."
    )
    st.caption("Límite de carga configurado: 500 MB")
    if _gemini_api_key():
        st.caption("👁️ Visión de fotografías: conectada")
    else:
        st.caption("👁️ Visión de fotografías: requiere configurar Gemini en Secrets")

    uploaded_file = st.file_uploader(
        "Seleccionar Excel",
        type=["xlsx", "xlsm"],
        key="tiara_excel_uploader",
        max_upload_size=500,
        accept_multiple_files=False,
        help=(
            "Carga el Excel original directamente desde Descargas. La aplicación "
            "lo lee internamente, sin que tengas que descomprimirlo. "
            "Se admiten archivos de hasta 500 MB en este cargador."
        ),
    )

    if uploaded_file is not None:
        file_bytes = uploaded_file.getvalue()
        signature = (
            uploaded_file.name,
            uploaded_file.size,
            hashlib.sha256(file_bytes).hexdigest(),
        )

        if st.session_state.uploaded_signature != signature:
            previous_data = st.session_state.uploaded_data

            try:
                with st.spinner(
                    "Procesando Excel… no necesitas descomprimirlo manualmente."
                ):
                    uploaded_file.seek(0)
                    new_data = build_from_uploaded_excel(uploaded_file)

                # El Excel válido se activa primero. La persistencia en
                # Supabase es secundaria: si la política de la tabla, la red o
                # la clave pública impiden guardar, la aplicación sigue usando
                # el Excel recién cargado durante esta sesión.
                st.session_state.uploaded_data = new_data
                supabase_save_error = None
                try:
                    save_supabase_data(new_data)
                except Exception as supabase_exc:
                    supabase_save_error = str(supabase_exc)

                st.session_state.uploaded_signature = signature
                st.session_state.last_question = ""
                st.session_state.last_answer = ""
                st.session_state.query_context = {"month": None, "year": None}

                meta = new_data.get("metadata", {})
                st.session_state.upload_message = (
                    f"✓ Excel activo: {meta.get('source_file', uploaded_file.name)} "
                    f"· {meta.get('source_size_mb', 0):.2f} MB"
                )
                if supabase_save_error:
                    st.session_state.upload_message += (
                        " · ⚠️ No se pudo guardar en Supabase; el Excel sí quedó activo."
                    )

            except Exception as exc:
                st.session_state.uploaded_data = previous_data
                st.session_state.upload_message = (
                    f"✕ No se pudo actualizar el Excel: {exc}"
                )

    if st.session_state.upload_message:
        st.caption(st.session_state.upload_message)

# Prioridad de fuentes:
# 1) Excel recién cargado en esta sesión.
# 2) Último snapshot persistente guardado en Supabase.
# 3) Base incluida en GitHub como respaldo.
data = st.session_state.uploaded_data
if not data:
    data = load_supabase_data()
if not data:
    data = load_data()

if not data:
    st.write("No se encontró la base de datos del Agente Tiara.")
    st.stop()

with st.sidebar:
    meta = data.get("metadata", {})
    st.markdown("### 📚 Base de datos")
    st.caption(f"Fuente activa: {meta.get('source_file', 'Excel de septiembre 2026')}")
    st.caption(f"Versión de la aplicación: {APP_VERSION}")
    updated_at = meta.get("source_updated_in_app")
    if updated_at:
        try:
            updated_dt = datetime.fromisoformat(str(updated_at).replace("Z", "+00:00"))
            if updated_dt.tzinfo is None:
                updated_dt = updated_dt.replace(tzinfo=APP_TIMEZONE)
            updated_text = updated_dt.astimezone(APP_TIMEZONE).strftime("%d/%m/%Y %H:%M:%S")
        except Exception:
            updated_text = str(updated_at)
    elif meta.get("supabase_created_at"):
        try:
            persisted_dt = datetime.fromisoformat(str(meta["supabase_created_at"]).replace("Z", "+00:00")).astimezone(APP_TIMEZONE)
            updated_text = persisted_dt.strftime("%d/%m/%Y %H:%M:%S")
        except Exception:
            updated_text = str(meta["supabase_created_at"])
    else:
        updated_text = APP_UPDATED
    st.caption(f"Actualización: {updated_text}")
    st.caption(f"Versión de la base: {meta.get('version', APP_VERSION)}")
    if meta.get("source_size_mb") is not None:
        st.caption(f"Tamaño procesado: {meta.get('source_size_mb')} MB")

# Estado de interfaz: solo mostramos la pregunta y respuesta ACTUALES.
# Se conserva internamente un contexto mínimo para entender frases como
# "ese mismo mes", sin conservar ni mostrar el historial de chat.
if "last_question" not in st.session_state:
    st.session_state.last_question = ""
if "last_answer" not in st.session_state:
    st.session_state.last_answer = ""
if "query_context" not in st.session_state:
    st.session_state.query_context = {"month": None, "year": None}
if "voice_audio_signature" not in st.session_state:
    st.session_state.voice_audio_signature = None

# ============================================================
# ENCABEZADO PRINCIPAL
# ============================================================
st.markdown("""
<div class="tiara-header">
    <div class="tiara-brand">
        <div class="tiara-brand-title">AGENTE <span>TIARA</span></div>
        <div class="tiara-brand-subtitle">TU ASISTENTE DE CONFIANZA</div>
    </div>
</div>
<div class="tiara-hero">
    <p>Centro inteligente de mantenimiento, operación, inventario, combustible y documentación del barco.</p>
</div>
""", unsafe_allow_html=True)

# Los datos administrativos quedan fuera de la pantalla principal.
# Se mantienen disponibles en la barra lateral.
hours = current_hours(data)
maintenance_n = len([
    r for r in maintenance_records(data)
    if norm(r["data"].get("mantenimiento")) not in ("motores", "generador")
])

with st.sidebar:
    st.markdown("### 📊 Datos de la base")
    st.write(f"**Hojas de datos:** {len(data.get('sheets', {}))}")
    st.write(
        f"**Registros:** {sum(len(data.get(k, [])) for k in ('maintenance_records', 'log_records', 'inventory_records', 'fuel_records', 'budget_records', 'invoice_records', 'cleaning_records', 'documents_records')):,}"
    )
    st.write(f"**Mantenimientos:** {maintenance_n}")
    st.divider()
    st.markdown("### ⏱️ Horómetros")
    for system, value in hours.items():
        st.write(
            f"**{system}:** {value} h"
            if value is not None
            else f"**{system}:** NO DETERMINADO"
        )

# ============================================================
# CONSULTAS RÁPIDAS — DISEÑO 2 + 2 + 1
# ============================================================
st.markdown('<div class="tiara-section-title">⚡ Consultas rápidas</div>', unsafe_allow_html=True)

quick = None

q1, q2 = st.columns(2, gap="small")
with q1:
    if st.button("📊  Análisis financiero", use_container_width=True):
        quick = "¿Cuáles son tus capacidades de análisis financiero?"
with q2:
    if st.button("🔧  Motores", use_container_width=True):
        quick = "¿Qué mantenimiento les toca próximamente a los motores?"

q3, q4 = st.columns(2, gap="small")
with q3:
    if st.button("🔩  Kit 300 h", use_container_width=True):
        quick = "¿Cuándo toca el próximo kit de 300 horas de los motores?"
with q4:
    if st.button("⚙️  Patas", use_container_width=True):
        quick = "¿Cuándo toca el próximo cambio de aceite de las patas?"

q5 = st.columns(1)[0]
with q5:
    if st.button("🔌  Generador", use_container_width=True):
        quick = "¿Cuándo toca el próximo cambio de aceite al generador?"

# ============================================================
# CHAT — TEXTO + VOZ + ENVIAR
# ============================================================
# Preferimos el chat nativo de Streamlit cuando la versión instalada soporta
# audio integrado. Así texto, micrófono y enviar son UN solo widget: no hay
# reproductor de audio suelto que se desborde en el teléfono.
# Si el Streamlit del despliegue es antiguo, conservamos el fallback anterior.

quick = quick or ""
voice_audio = None
typed_question = ""
send_question = False

try:
    _chat_input_params = inspect.signature(st.chat_input).parameters
    _supports_audio_chat = "accept_audio" in _chat_input_params
except Exception:
    _supports_audio_chat = False

if _supports_audio_chat:
    _chat_value = st.chat_input(
        "Escribe tu consulta o habla…",
        key="tiara_chat_input",
        accept_audio=True,
        audio_sample_rate=16000,
    )

    if _chat_value is not None:
        # Streamlit entrega ChatInputValue, que permite acceso por atributo
        # y también como diccionario. Usamos ambos para máxima compatibilidad
        # entre versiones del despliegue sin modificar el widget visual.
        typed_question = (
            getattr(_chat_value, "text", None)
            if not isinstance(_chat_value, str)
            else _chat_value
        ) or ""
        typed_question = str(typed_question).strip()

        voice_audio = getattr(_chat_value, "audio", None)
        if voice_audio is None and isinstance(_chat_value, dict):
            voice_audio = _chat_value.get("audio")

        send_question = bool(typed_question or voice_audio is not None)
else:
    # Fallback para versiones antiguas de Streamlit.
    with st.container(key="tiara_composer"):
        c_voice, c_text, c_send = st.columns([1, 7, 1], gap="small", vertical_alignment="center")

        with c_voice:
            voice_audio = st.audio_input(
                "🎙️",
                key="tiara_voice_input",
                label_visibility="collapsed",
            )

        with c_text:
            typed_question = st.text_input(
                "Pregunta",
                placeholder="Escribe tu consulta…",
                label_visibility="collapsed",
                key="tiara_text_input",
            )

        with c_send:
            send_question = st.button(
                "➤",
                use_container_width=True,
                help="Enviar la pregunta al Agente Tiara",
                key="tiara_send_question",
            )

        st.markdown(
            '<div class="tiara-help">Puedes escribir o hablar. Pulsa ➤ para enviar.</div>',
            unsafe_allow_html=True,
        )

# La interacción manual y la de voz convergen aquí: solo después de tener
# una pregunta válida se llama al mismo run_agent().
if send_question:
    typed_question = (typed_question or "").strip()

    if typed_question:
        quick = typed_question
    elif voice_audio is not None:
        try:
            audio_bytes = voice_audio.getvalue()
        except Exception:
            audio_bytes = bytes(voice_audio) if voice_audio else b""

        audio_type = getattr(voice_audio, "type", None) or "audio/wav"
        audio_signature = hashlib.sha256(audio_bytes).hexdigest() if audio_bytes else None

        if audio_signature != st.session_state.voice_audio_signature:
            try:
                with st.spinner("🎙️ Procesando tu pregunta…"):
                    voice_question, voice_error = _gemini_audio_transcribe(audio_bytes, audio_type)
            except Exception as exc:
                voice_question, voice_error = None, f"ERROR_TRANSCRIPCION: {exc}"

            if voice_question:
                st.session_state.voice_audio_signature = audio_signature
                quick = voice_question.strip()
            else:
                # Mostrar el error en el mismo campo de respuesta para que el
                # usuario nunca quede sin saber qué falló en la ruta de voz.
                st.session_state.last_question = "🎙️ Pregunta por voz"
                st.session_state.last_answer = (
                    "No pude convertir el audio en una pregunta de texto. "
                    f"Detalle: {voice_error}"
                )

if quick:
    # Ejecutar con el contexto oculto de la consulta inmediatamente anterior.
    st.session_state.last_question = quick
    st.session_state.last_answer = run_agent(
        data, quick, st.session_state.query_context
    )

    # Actualizar SOLO el contexto útil (mes/año). No se guarda la pregunta ni
    # la respuesta anterior como historial.
    new_thought = think(quick)
    if new_thought.get("month") is not None:
        st.session_state.query_context["month"] = new_thought["month"]
    if new_thought.get("year") is not None:
        st.session_state.query_context["year"] = new_thought["year"]

# ============================================================
# RESULTADO ACTUAL — CAMPO DE RESPUESTA
# ============================================================
# La respuesta del agente siempre se muestra escrita en un único campo
# dedicado. Esto es igual para preguntas escritas y preguntas hechas por voz:
# la voz se convierte primero a texto y luego pasa por run_agent().
if st.session_state.last_question:
    st.markdown(
        '<div class="tiara-response-label">RESPUESTA DEL AGENTE TIARA</div>',
        unsafe_allow_html=True,
    )

    with st.container(key="tiara_response"):
        st.markdown(st.session_state.last_answer)

# La fuente de datos queda disponible en la barra lateral.
